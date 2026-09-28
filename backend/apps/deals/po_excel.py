from __future__ import annotations

import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from apps.deals.management.commands.seed_funnel import parse_value
from apps.deals.po_constants import (
    AGREEMENTS,
    BILLING_CYCLES,
    BILLING_DUES,
    BILLING_ENTITIES,
    PO_SERVICE_PILLARS,
    PO_WORKFLOW_HEADERS,
)

HEADER_MARKERS = {'customer', 'billing entity', 'po number', 'po workflow'}


def cell_text(value, *, keep_dash: bool = False) -> str:
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if text == '*':
        return ''
    if keep_dash:
        return text
    if text in {'', '-'}:
        return ''
    return text


def norm_header(value) -> str:
    text = cell_text(value).lower().replace('₹', ' ')
    text = re.sub(r'[^a-z0-9]+', ' ', text)
    return ' '.join(text.split())


def header_index(row) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, value in enumerate(row):
        key = norm_header(value)
        if key:
            mapping[key] = idx
    return mapping


def cell_at(row, headers: dict[str, int], *names):
    for name in names:
        idx = headers.get(norm_header(name))
        if idx is not None and idx < len(row):
            return row[idx]
    return None


def as_date(value) -> date | None:
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = cell_text(value)
    if not text or text in {'-', 'NA', 'na'}:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def as_money(value) -> Decimal | None:
    if value is None or value == '':
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = cell_text(value)
    if not text:
        return None
    parsed, _currency = parse_value(text)
    if parsed is not None:
        return parsed
    try:
        return Decimal(text.replace(',', ''))
    except (InvalidOperation, ValueError):
        return None


def is_skip_row(values: list[str]) -> bool:
    if not any(values):
        return True
    if any(value.lower() in HEADER_MARKERS for value in values if value):
        return True
    if all(value in {'*', ''} for value in values):
        return True
    return False


def parse_po_row(raw, headers: dict[str, int]) -> dict | None:
    customer = cell_text(cell_at(raw, headers, 'customer', 'customer name', 'company'))
    billing_entity = cell_text(cell_at(raw, headers, 'billing entity'))
    service_pillar = cell_text(cell_at(raw, headers, 'service pillar'))
    po_number = cell_text(cell_at(raw, headers, 'po number'), keep_dash=True)
    if is_skip_row([customer, billing_entity, service_pillar, po_number]):
        return None
    return {
        'customer_name': customer[:255],
        'billing_entity': billing_entity[:120],
        'service_pillar': service_pillar[:120],
        'po_number': po_number[:80],
        'po_date': as_date(cell_at(raw, headers, 'po date')),
        'agreement': cell_text(cell_at(raw, headers, 'agreement'))[:40],
        'start_date': as_date(cell_at(raw, headers, 'start date')),
        'end_date': as_date(cell_at(raw, headers, 'end date')),
        'expiry': as_date(cell_at(raw, headers, 'expiry')),
        'total_value': as_money(cell_at(raw, headers, 'total value')),
        'value_per_invoice': as_money(cell_at(raw, headers, 'value per invoice')),
        'billing_cycle': cell_text(cell_at(raw, headers, 'billing cycle'))[:80],
        'billing_due': cell_text(cell_at(raw, headers, 'billing due'))[:120],
    }


def load_po_rows_from_workbook(workbook) -> list[tuple[int, dict]]:
    sheet = workbook['PO_Workflow'] if 'PO_Workflow' in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    iterator = sheet.iter_rows(min_row=1, max_col=20, values_only=True)
    try:
        headers = header_index(next(iterator))
    except StopIteration:
        return []
    rows: list[tuple[int, dict]] = []
    for excel_row, raw in enumerate(iterator, start=2):
        parsed = parse_po_row(raw, headers)
        if parsed:
            rows.append((excel_row, parsed))
    return rows


def _autosize(sheet: Worksheet, width: int = 20):
    for index in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _add_list_validation(sheet: Worksheet, col: int, options: list[str], start_row: int = 3):
    formula = '"' + ','.join(options) + '"'
    if len(formula) > 250:
        return
    validation = DataValidation(type='list', formula1=formula, allow_blank=True)
    letter = get_column_letter(col)
    validation.add(f'{letter}{start_row}:{letter}200')
    sheet.add_data_validation(validation)


def build_po_workbook(rows: list[list], *, with_star_row: bool = False, with_validation: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'PO_Workflow'
    sheet.append(PO_WORKFLOW_HEADERS)
    if with_star_row:
        sheet.append(['*'] * len(PO_WORKFLOW_HEADERS))
    for row in rows:
        sheet.append(row)
    if with_validation:
        _add_list_validation(sheet, 2, BILLING_ENTITIES)
        _add_list_validation(sheet, 3, PO_SERVICE_PILLARS)
        _add_list_validation(sheet, 6, AGREEMENTS)
        _add_list_validation(sheet, 12, BILLING_CYCLES)
        _add_list_validation(sheet, 13, BILLING_DUES)
    _autosize(sheet)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def open_upload(content: bytes):
    return load_workbook(io.BytesIO(content), data_only=True)


def po_export_row(item) -> list:
    return [
        item.customer_name,
        item.billing_entity,
        item.service_pillar,
        item.po_number,
        item.po_date.isoformat() if item.po_date else '',
        item.agreement,
        item.start_date.isoformat() if item.start_date else '',
        item.end_date.isoformat() if item.end_date else '',
        item.expiry.isoformat() if item.expiry else '',
        float(item.total_value) if item.total_value is not None else '',
        float(item.value_per_invoice) if item.value_per_invoice is not None else '',
        item.billing_cycle,
        item.billing_due,
    ]


def dummy_po_rows() -> list[list]:
    return [[
        'Chavans',
        'India',
        'BizOps',
        'PO-2526-0048',
        '2025-07-14',
        'NA',
        '2025-08-01',
        '2026-03-31',
        '2026-08-26',
        17100000,
        300000,
        'Monthly',
        '1st of Month',
    ]]
