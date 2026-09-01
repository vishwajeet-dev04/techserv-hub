from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.business.models import Company, CompanyStatus
from apps.deals.models import Deal, DealStage
from apps.deals.services.deal_service import set_default_probability
from apps.users.models import User, UserRole

DEFAULT_PASSWORD = 'TechServ2026!'
AED_TO_INR = Decimal('23.5')

TEAM = {
    'AC': ('AC', 'ac@techserv.ai'),
    'PS': ('PS', 'ps@techserv.ai'),
    'SM': ('SM', 'sm@techserv.ai'),
    'DS': ('DS', 'ds@techserv.ai'),
    'AB': ('AB', 'ab@techserv.ai'),
    'TS': ('TS', 'ts@techserv.ai'),
    'ASIF': ('Asif', 'asif@techserv.ai'),
    'PRAVEEN': ('Praveen', 'praveen@techserv.ai'),
    'SAJI': ('Saji', 'saji@techserv.ai'),
}

PARTNER_NAMES = {
    'redington', 'savex', 'galaxy', 'chavans', 'chavans technologies',
    'inflow technologies', 'sunfinity', 'ptc, sg', 'perfekt, au',
    'concept, pune', 'technobind', 'wysetek', 'ashtech', 'pentagon',
    'veeam', 'lenovo enterprise', 'lenovo isg', 'hpi singapore',
    'commvault', 'nutanix', 'netapp', 'commvault - marketing', 'commvault (gsi)',
}

WON_KEYWORDS = ('po received', 'kickoff', 'kick off', 'closed - costing')
LOST_KEYWORDS = ('lost', 'drop', 'cancelled', 'canceled', 'refund')
NEGOTIATION_KEYWORDS = (
    'awaiting po', 'pi sent', 'contracting', 'commercial closure',
    'costing discussion', 'po etc', 'agreements done', 'formalities',
)
PROPOSAL_KEYWORDS = ('proposal', 'sow', 'commercial')
QUALIFIED_KEYWORDS = (
    'workshop', 'demo', 'engaged', 'nda', 'prototype', 'architecture',
    'sample data', 'alignment call',
)


def _cell(value) -> str:
    if value is None:
        return ''
    text = str(value).strip()
    if text in {'None', '`', '-', 'TBD', 'tbd'}:
        return ''
    return text


def _owner_codes(raw: str) -> list[str]:
    if not raw:
        return []
    parts = re.split(r'[/&,]| and ', raw)
    codes = []
    for part in parts:
        token = part.strip().upper()
        if not token:
            continue
        if token in TEAM:
            codes.append(token)
        elif token.replace(' ', '') in TEAM:
            codes.append(token.replace(' ', ''))
    return codes


def parse_value(raw: str) -> tuple[Decimal, str]:
    text = _cell(raw)
    if not text:
        return Decimal('0'), 'INR'

    monthly = '/month' in text.lower() or 'per month' in text.lower()
    currency = 'AED' if 'aed' in text.lower() else 'INR'
    cleaned = re.sub(r'[^0-9.]', '', text.replace(',', ''))
    if not cleaned:
        return Decimal('0'), currency
    amount = Decimal(cleaned)
    if monthly:
        amount *= 12
    if currency == 'AED':
        amount *= AED_TO_INR
        currency = 'INR'
    return amount, currency


def parse_probability(raw, stage: str) -> int:
    text = _cell(raw)
    if not text:
        return set_default_probability(stage)
    try:
        value = float(text)
    except ValueError:
        return set_default_probability(stage)
    if value <= 1:
        value *= 100
    return max(0, min(100, int(round(value))))


def map_stage(status: str, remarks: str = '') -> str:
    blob = f'{status} {remarks}'.lower()
    if any(keyword in blob for keyword in LOST_KEYWORDS):
        return DealStage.LOST
    if any(keyword in blob for keyword in WON_KEYWORDS):
        return DealStage.WON
    if 'closed' in blob and 'pending' not in blob:
        return DealStage.WON
    if any(keyword in blob for keyword in NEGOTIATION_KEYWORDS):
        return DealStage.NEGOTIATION
    if any(keyword in blob for keyword in PROPOSAL_KEYWORDS):
        return DealStage.PROPOSAL
    if any(keyword in blob for keyword in QUALIFIED_KEYWORDS):
        return DealStage.QUALIFIED
    if any(keyword in blob for keyword in ('requirement', 'sow to be defined')):
        return DealStage.REQUIREMENT
    return DealStage.LEAD


def company_status_for(name: str, stages: list[str]) -> str:
    if DealStage.WON in stages:
        return CompanyStatus.CUSTOMER
    if name.strip().lower() in PARTNER_NAMES:
        return CompanyStatus.PARTNER
    return CompanyStatus.PROSPECT


def default_xlsx_path() -> Path:
    return Path(settings.BASE_DIR).parent / 'TechServ Funnel Sheet_Updated 28062026.xlsx'


def load_workbook_rows(path: Path) -> list[dict]:
    try:
        import openpyxl
    except ImportError as exc:
        raise CommandError('openpyxl is required to read the funnel sheet. pip install openpyxl') from exc

    workbook = openpyxl.load_workbook(path, data_only=True)
    rows: list[dict] = []

    services = workbook['Services and Product Pipeline']
    for index, row in enumerate(services.iter_rows(min_row=3, max_row=32, max_col=8, values_only=True), start=3):
        customer = _cell(row[0])
        if not customer or customer.lower() in {'customer', 'it channel partners:'}:
            continue
        rows.append({
            'company': customer,
            'pillar': _cell(row[1]) or 'Services',
            'sub_category': _cell(row[2]),
            'owner': _cell(row[3]),
            'status': _cell(row[4]),
            'value': _cell(row[5]),
            'probability': row[6],
            'remarks': _cell(row[7]),
            'source': 'Services & Product',
            'sheet_row': index,
        })

    for index, row in enumerate(services.iter_rows(min_row=46, max_row=56, max_col=8, values_only=True), start=46):
        customer = _cell(row[0])
        if not customer or customer.lower() == 'customer':
            continue
        rows.append({
            'company': customer,
            'pillar': 'Sprouts',
            'sub_category': 'Channel partner',
            'owner': _cell(row[1]),
            'status': _cell(row[3]),
            'value': '',
            'probability': None,
            'remarks': _cell(row[7]),
            'source': 'Services & Product',
            'sheet_row': index,
        })

    software = workbook['Software & AI Dev Pipeline']
    lost_section = False
    for index, row in enumerate(software.iter_rows(min_row=4, max_row=25, max_col=9, values_only=True), start=4):
        customer = _cell(row[0])
        if not customer:
            continue
        if customer.lower() == 'lost':
            lost_section = True
            continue
        status = _cell(row[4])
        remarks = _cell(row[8])
        if lost_section:
            status = status or 'Lost'
            remarks = remarks or 'Lost'
        rows.append({
            'company': customer,
            'pillar': _cell(row[1]) or 'Software Development',
            'sub_category': _cell(row[2]),
            'owner': _cell(row[3]),
            'status': status,
            'value': _cell(row[5]),
            'probability': row[6],
            'remarks': remarks,
            'source': 'Software & AI',
            'sheet_row': index,
        })
    return rows


class Command(BaseCommand):
    help = 'Load companies and deals from the TechServ funnel Excel sheet.'

    def add_arguments(self, parser):
        parser.add_argument('--path', type=str, default='', help='Path to the funnel xlsx file')
        parser.add_argument('--replace', action='store_true', help='Replace previously seeded funnel deals')

    def handle(self, *args, **options):
        path = Path(options['path']) if options['path'] else default_xlsx_path()
        if not path.exists():
            raise CommandError(f'Funnel sheet not found: {path}')

        rows = load_workbook_rows(path)
        if not rows:
            raise CommandError('No funnel rows were parsed from the spreadsheet.')

        users = self._ensure_users()
        created_companies = 0
        created_deals = 0
        updated_deals = 0

        with transaction.atomic():
            if options['replace']:
                deleted, _ = Deal.objects.filter(notes__contains='[funnel-sheet]').delete()
                self.stdout.write(f'Removed {deleted} previously seeded deals.')

            company_stages: dict[str, list[str]] = {}
            for row in rows:
                key = row['company'].strip()
                stage = map_stage(row['status'], row['remarks'])
                company_stages.setdefault(key, []).append(stage)

            companies: dict[str, Company] = {}
            for name, stages in company_stages.items():
                status = company_status_for(name, stages)
                company, was_created = Company.objects.get_or_create(
                    name=name,
                    defaults={
                        'status': status,
                        'company_type': 'Customer',
                        'industry': 'Technology',
                        'notes': 'Imported from TechServ Funnel Sheet (28 Jun 2026).',
                    },
                )
                if was_created:
                    created_companies += 1
                elif company.status != status and status == CompanyStatus.CUSTOMER:
                    company.status = status
                    company.save(update_fields=['status', 'updated_at'])
                companies[name] = company

            for row in rows:
                company = companies[row['company'].strip()]
                codes = _owner_codes(row['owner'])
                owner = users.get(codes[0]) if codes else None
                if owner and company.owner_id is None:
                    company.owner = owner
                    company.save(update_fields=['owner', 'updated_at'])

                stage = map_stage(row['status'], row['remarks'])
                value, currency = parse_value(row['value'])
                probability = parse_probability(row['probability'], stage)
                pillar = row['pillar']
                sub_category = row['sub_category']
                deal_name = self._deal_name(company.name, pillar, sub_category)
                notes = self._notes(row)

                defaults = {
                    'company': company,
                    'description': row['status'],
                    'value': value,
                    'currency': currency,
                    'stage': stage,
                    'probability': probability,
                    'owner': owner,
                    'source': row['source'],
                    'service_pillar': pillar,
                    'sub_category': sub_category,
                    'notes': notes,
                }
                deal, was_created = Deal.objects.update_or_create(
                    name=deal_name,
                    company=company,
                    defaults=defaults,
                )
                if was_created:
                    created_deals += 1
                else:
                    updated_deals += 1

        self.stdout.write(self.style.SUCCESS(
            f'Seeded funnel sheet: {created_companies} companies, '
            f'{created_deals} deals created, {updated_deals} deals updated.'
        ))

    def _ensure_users(self) -> dict[str, User]:
        users = {}
        for code, (name, email) in TEAM.items():
            user, created = User.objects.get_or_create(
                email=email,
                defaults={
                    'name': name,
                    'role': UserRole.SALES,
                },
            )
            if created:
                user.set_password(DEFAULT_PASSWORD)
                user.save(update_fields=['password'])
                self.stdout.write(f'Created user {name} <{email}>')
            users[code] = user
        return users

    def _deal_name(self, company: str, pillar: str, sub_category: str) -> str:
        if sub_category:
            return f'{company} · {pillar} · {sub_category}'[:255]
        if pillar:
            return f'{company} · {pillar}'[:255]
        return company

    def _notes(self, row: dict) -> str:
        parts = ['[funnel-sheet]']
        if row['status']:
            parts.append(f"Status: {row['status']}")
        if row['remarks']:
            parts.append(row['remarks'])
        if row['owner']:
            parts.append(f"Owners: {row['owner']}")
        return '\n\n'.join(parts)
