import uuid
from datetime import date

from rest_framework import status
from rest_framework.test import APITestCase

from apps.business.models import Company, CompanyStatus
from apps.deals.models import Deal, DealActivity, DealActivityType, DealStage
from apps.deals.services.deal_service import get_deal, set_default_probability
from apps.users.models import User, UserRole


class DealDetailAPITests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            email='sales@test.com',
            name='Sales User',
            password='pass12345',
            role=UserRole.SALES,
        )
        self.client.force_authenticate(self.user)
        self.company = Company.objects.create(
            name='Acme Corp',
            status=CompanyStatus.CUSTOMER,
            owner=self.user,
        )
        self.deal = Deal.objects.create(
            company=self.company,
            name='ERP rollout',
            description='Implementation',
            value=1500000,
            stage=DealStage.PROPOSAL,
            probability=60,
            owner=self.user,
            source='Referral',
            notes='Priority account',
            expected_close_date=date(2026, 9, 30),
        )
        self.activity = DealActivity.objects.create(
            deal=self.deal,
            type=DealActivityType.CALL,
            subject='Discovery call',
            description='Covered requirements',
            activity_date=date(2026, 8, 1),
            created_by=self.user,
        )

    def test_get_deal_service_includes_company_and_activities(self):
        deal = get_deal(self.deal.id)
        self.assertEqual(deal.company.name, 'Acme Corp')
        self.assertEqual(deal.activities.count(), 1)

    def test_retrieve_deal_includes_activities_and_weighted_value(self):
        response = self.client.get(f'/api/v1/deals/{self.deal.id}/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['name'], 'ERP rollout')
        self.assertEqual(response.data['company_name'], 'Acme Corp')
        self.assertEqual(response.data['stage'], DealStage.PROPOSAL)
        self.assertEqual(float(response.data['weighted_value']), 900000.0)
        self.assertEqual(len(response.data['activities']), 1)
        self.assertEqual(response.data['activities'][0]['subject'], 'Discovery call')

    def test_retrieve_missing_deal_returns_404(self):
        response = self.client.get(f'/api/v1/deals/{uuid.uuid4()}/')
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assertTrue(response.data['error'])

    def test_patch_stage_sets_default_probability(self):
        response = self.client.patch(
            f'/api/v1/deals/{self.deal.id}/',
            {'stage': DealStage.NEGOTIATION},
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['stage'], DealStage.NEGOTIATION)
        self.assertEqual(response.data['probability'], set_default_probability(DealStage.NEGOTIATION))
        self.deal.refresh_from_db()
        self.assertEqual(self.deal.stage, DealStage.NEGOTIATION)

    def test_list_activities_for_deal(self):
        response = self.client.get(f'/api/v1/deals/{self.deal.id}/activities/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        results = response.data['results'] if isinstance(response.data, dict) else response.data
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]['type'], DealActivityType.CALL)

    def test_create_activity(self):
        response = self.client.post(
            f'/api/v1/deals/{self.deal.id}/activities/',
            {
                'type': DealActivityType.MEETING,
                'subject': 'Proposal review',
                'description': 'Walked through scope',
                'activity_date': '2026-08-19',
            },
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['subject'], 'Proposal review')
        self.assertEqual(response.data['created_by_name'], 'Sales User')
        self.assertTrue(
            DealActivity.objects.filter(deal=self.deal, subject='Proposal review').exists()
        )

    def test_unauthenticated_deal_detail_is_rejected(self):
        self.client.force_authenticate(user=None)
        response = self.client.get(f'/api/v1/deals/{self.deal.id}/')
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_pipeline_summary_includes_pillar_and_funnel(self):
        self.deal.service_pillar = 'Marketing'
        self.deal.sub_category = 'ICP/ABM'
        self.deal.source = 'Services & Product'
        self.deal.save()
        Deal.objects.create(
            company=self.company,
            name='Won deal',
            value=500000,
            stage=DealStage.WON,
            probability=100,
            owner=self.user,
            service_pillar='Consulting',
            source='Services & Product',
        )
        Deal.objects.create(
            company=self.company,
            name='Lost deal',
            value=200000,
            stage=DealStage.LOST,
            probability=0,
            owner=self.user,
            service_pillar='Marketing',
            source='Software & AI',
        )
        response = self.client.get('/api/v1/deals/pipeline/summary/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.data
        self.assertEqual(data['active_deals_count'], 1)
        self.assertEqual(data['won_count'], 1)
        self.assertEqual(data['lost_count'], 1)
        self.assertEqual(data['win_rate'], 50.0)
        pillars = {row['name']: row for row in data['by_pillar']}
        self.assertEqual(pillars['Marketing']['count'], 2)
        self.assertEqual(pillars['Consulting']['won_count'], 1)
        funnel_stages = {row['stage'] for row in data['funnel']}
        self.assertIn(DealStage.PROPOSAL, funnel_stages)
        self.assertIn(DealStage.WON, funnel_stages)
        self.assertIn(DealStage.LOST, funnel_stages)

    def test_list_deals_filter_by_service_pillar(self):
        self.deal.service_pillar = 'Marketing'
        self.deal.save()
        response = self.client.get('/api/v1/deals/', {'service_pillar': 'Marketing'})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        names = [item['name'] for item in response.data['results']]
        self.assertEqual(names, ['ERP rollout'])
        self.assertEqual(response.data['results'][0]['service_pillar'], 'Marketing')


class PoWorkflowAPITests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            email='sales-po@test.com',
            name='Sales User',
            password='pass12345',
            role=UserRole.SALES,
        )
        self.client.force_authenticate(self.user)

    def test_create_requires_starred_fields(self):
        response = self.client.post('/api/v1/deals/po-workflow/', {'customer_name': 'Chavans'}, format='json')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_create_saves_dash_po_number(self):
        response = self.client.post(
            '/api/v1/deals/po-workflow/',
            {
                'customer_name': 'New Era',
                'billing_entity': 'India',
                'service_pillar': 'AIDA',
                'po_number': '-',
                'agreement': 'Yes',
                'billing_cycle': 'Monthly',
                'billing_due': 'Event Specific',
                'total_value': '200000',
            },
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(response.data['po_number'], '-')
        self.assertEqual(response.data['customer_name'], 'New Era')
        listed = self.client.get('/api/v1/deals/po-workflow/', {'search': 'New Era'})
        self.assertEqual(listed.data['count'], 1)
        self.assertEqual(listed.data['results'][0]['po_display'], 'PO-New Era')
        self.assertEqual(listed.data['results'][0]['status'], 'Active')

    def test_import_skips_star_and_empty_keeps_dash(self):
        from io import BytesIO

        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook, load_workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'PO_Workflow'
        sheet.append([
            'Customer', 'Billing Entity', 'Service Pillar', 'PO Number', 'PO Date', 'Agreement',
            'Start Date', 'End Date', 'Expiry', 'Total value', 'Value per Invoice', 'Billing Cycle', 'Billing Due',
        ])
        sheet.append(['*'] * 13)
        sheet.append([''] * 13)
        sheet.append([
            'New Era', 'India', 'AIDA', '-', '-', 'Yes',
            None, None, None, 200000, 50000, 'Monthly', 'Event Specific',
        ])
        buffer = BytesIO()
        workbook.save(buffer)
        uploaded = SimpleUploadedFile('po.xlsx', buffer.getvalue())
        result = self.client.post('/api/v1/deals/po-workflow/import/', {'file': uploaded}, format='multipart')
        self.assertEqual(result.status_code, status.HTTP_200_OK, result.data)
        self.assertEqual(result.data['created'], 1)
        self.assertEqual(result.data['errors'], [])
        from apps.deals.models import PoWorkflow
        row = PoWorkflow.objects.get()
        self.assertEqual(row.po_number, '-')
        self.assertEqual(row.customer_name, 'New Era')
        template = self.client.get('/api/v1/deals/po-workflow/template/')
        self.assertEqual(template.status_code, status.HTTP_200_OK)
        book = load_workbook(BytesIO(template.content))
        star = [cell.value for cell in next(book.active.iter_rows(min_row=2, max_row=2))]
        self.assertTrue(all(value == '*' for value in star[:13]))
        exported = self.client.get('/api/v1/deals/po-workflow/export/')
        self.assertEqual(exported.status_code, status.HTTP_200_OK)

    def test_viewer_cannot_read_po_workflow(self):
        viewer = User.objects.create_user(
            email='view-po@test.com',
            name='Viewer',
            password='pass12345',
            role=UserRole.VIEWER,
        )
        self.client.force_authenticate(viewer)
        response = self.client.get('/api/v1/deals/po-workflow/')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        dash = self.client.get('/api/v1/deals/po-workflow/dashboard/')
        self.assertEqual(dash.status_code, status.HTTP_403_FORBIDDEN)

    def test_dashboard_filters_and_status(self):
        from datetime import date, timedelta

        from apps.deals.models import PoWorkflow
        from apps.business.models import Company

        company = Company.objects.create(name='Chavans')
        today = date.today()
        PoWorkflow.objects.create(
            company=company,
            customer_name='Chavans',
            billing_entity='India',
            service_pillar='BizOps',
            po_number='PO-2526-0048',
            agreement='NA',
            start_date=today - timedelta(days=400),
            end_date=today - timedelta(days=10),
            total_value=17100000,
            value_per_invoice=300000,
            billing_cycle='Monthly',
            billing_due='1st of Month',
        )
        PoWorkflow.objects.create(
            company=company,
            customer_name='New Era',
            billing_entity='SG',
            service_pillar='AIDA',
            po_number='-',
            agreement='Yes',
            total_value=200000,
            value_per_invoice=50000,
            billing_cycle='Monthly',
            billing_due='Event Specific',
        )
        dash = self.client.get('/api/v1/deals/po-workflow/dashboard/')
        self.assertEqual(dash.status_code, status.HTTP_200_OK, dash.data)
        self.assertEqual(dash.data['kpis']['active_pos'], 1)
        self.assertEqual(dash.data['kpis']['total_contract_value'], 17300000.0)
        displays = {row['po_display'] for row in dash.data['rows']}
        self.assertIn('PO-New Era', displays)
        statuses = {row['customer_name']: row['status'] for row in dash.data['rows']}
        self.assertEqual(statuses['Chavans'], 'Completed')
        self.assertEqual(statuses['New Era'], 'Active')
        india = self.client.get('/api/v1/deals/po-workflow/dashboard/', {'billing_entity': 'India'})
        self.assertEqual(len(india.data['rows']), 1)
        client = self.client.get('/api/v1/deals/po-workflow/dashboard/', {'customer': 'New Era'})
        self.assertEqual(len(client.data['rows']), 1)
        far = self.client.get('/api/v1/deals/po-workflow/dashboard/', {'from': '2099-01-01', 'to': '2099-01-31'})
        self.assertEqual(len(far.data['rows']), 0)
        overlap = self.client.get(
            '/api/v1/deals/po-workflow/dashboard/',
            {'from': (today - timedelta(days=20)).isoformat(), 'to': today.isoformat()},
        )
        names = {row['customer_name'] for row in overlap.data['rows']}
        self.assertEqual(names, {'Chavans'})


class FunnelSheetMappingTests(APITestCase):
    def test_status_maps_to_pipeline_stages(self):
        from apps.deals.management.commands.seed_funnel import map_stage, parse_probability, parse_value

        self.assertEqual(map_stage('Closed - PO received but project not executed'), DealStage.WON)
        self.assertEqual(map_stage('Lost to Wizard.io'), DealStage.LOST)
        self.assertEqual(map_stage('Proposal sent to Kush'), DealStage.PROPOSAL)
        self.assertEqual(map_stage('KickOff Today'), DealStage.WON)
        self.assertEqual(map_stage('Warm - Discussions Open'), DealStage.LEAD)

        value, currency = parse_value('350000/month')
        self.assertEqual(value, 4200000)
        self.assertEqual(currency, 'INR')
        self.assertEqual(parse_probability(0.7, DealStage.PROPOSAL), 70)


