from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError as DrfValidationError

from apps.pulse.excel import find_sheet, load_contacts_from_workbook, load_opportunities_from_workbook
from apps.pulse.models import Opportunity, PulseContact, Reminder
from apps.pulse.services import pulse_service


def default_xlsx_path() -> Path:
    repo = Path(settings.BASE_DIR).parent
    preferred = repo / 'PROJECT_CONTEXT' / 'Files' / 'Sales Pipeline v2.xlsx'
    if preferred.exists():
        return preferred
    return repo / 'TechServ Funnel Sheet_Updated 28062026.xlsx'


class Command(BaseCommand):
    help = 'Load Pulse CRM contacts and opportunities from Sales Pipeline Excel (Contact + Lead - Deal sheets).'

    def add_arguments(self, parser):
        parser.add_argument('--file', type=str, default='', help='Path to xlsx')
        parser.add_argument(
            '--reset',
            action='store_true',
            help='Delete existing Pulse contacts, opportunities, and reminders before loading',
        )

    def handle(self, *args, **options):
        path = Path(options['file']) if options['file'] else default_xlsx_path()
        if not path.exists():
            raise CommandError(f'Excel file not found: {path}')
        workbook = load_workbook(path, data_only=True)
        if find_sheet(workbook, 'contact') is None and find_sheet(workbook, 'lead') is None:
            self.stdout.write(self.style.WARNING('Workbook sheets: %s' % workbook.sheetnames))
        contact_rows = load_contacts_from_workbook(workbook)
        opportunity_rows = load_opportunities_from_workbook(workbook)
        if not contact_rows and not opportunity_rows:
            raise CommandError('No Contact or Lead - Deal rows found in workbook.')

        created_contacts = 0
        created_opps = 0
        skipped_opps = 0
        with transaction.atomic():
            if options['reset']:
                Reminder.objects.all().delete()
                Opportunity.objects.all().delete()
                PulseContact.objects.all().delete()
            for _row, payload in contact_rows:
                created_at = payload.pop('created_at', None)
                pulse_service.create_contact(strict=False, created_at=created_at, **payload)
                created_contacts += 1
            for _row, payload in opportunity_rows:
                try:
                    pulse_service.create_opportunity(strict=False, **payload)
                    created_opps += 1
                except DrfValidationError:
                    skipped_opps += 1

        self.stdout.write(self.style.SUCCESS(
            f'Seeded {created_contacts} Pulse contacts and {created_opps} opportunities '
            f'({skipped_opps} opportunities skipped: no matching contact) from {path}'
        ))
