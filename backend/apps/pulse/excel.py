from __future__ import annotations

import io
import re
from datetime import date, datetime
from decimal import Decimal

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from apps.deals.management.commands.seed_funnel import parse_value
from apps.pulse.constants import (
    CONTACT_HEADERS,
    CURRENCIES,
    OPPORTUNITY_HEADERS,
    SALES_STAGES,
    SERVICE_PILLARS,
    SUB_CATEGORIES,
    coerce_currency,
    coerce_pillar,
    coerce_sales_stage,
    coerce_sub_category,
)

SKIP_MARKERS = {'', '*', '-', 'none', 'n/a', 'na'}
HEADER_MARKERS = {
    'company name', 'contact name', 'name', 'customer', 'opportunity',
    'prospect number', 'prospect name',
}


def cell_text(value) -> str:
    if value is None:
        return ''
    text = str(value).strip()
    if text.lower() in SKIP_MARKERS:
        return ''
    return text


def norm_header(value) -> str:
    text = cell_text(value).lower().replace('₹', ' ')
    text = re.sub(r'[^a-z0-9]+', ' ', text)
    return ' '.join(text.split())


def header_index(row) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, value in enumerate(row):
        key = norm_header(value)
        if key:
            mapping[key] = idx
    return mapping


def cell_at(row, headers: dict[str, int], *names):
    for name in names:
        idx = headers.get(norm_header(name))
        if idx is not None and idx < len(row):
            return row[idx]
    return None


_CALLING_CODES = tuple(sorted({
    '1', '7', '20', '27', '30', '31', '32', '33', '34', '39', '41', '43', '44', '45', '46', '47', '48', '49',
    '51', '52', '54', '55', '56', '57', '58', '60', '61', '62', '63', '64', '65', '66', '81', '82', '84', '86',
    '90', '91', '92', '93', '94', '95', '98', '212', '213', '216', '234', '254', '255', '256',
    '351', '352', '353', '354', '355', '356', '357', '358', '359', '370', '371', '372', '373', '374', '375',
    '380', '381', '382', '385', '386', '387', '420', '421', '423', '852', '853', '855', '856', '880', '886',
    '960', '961', '962', '963', '964', '965', '966', '967', '968', '970', '971', '972', '973', '974', '975',
    '976', '977', '992', '993', '994', '995', '996', '998',
}, key=len, reverse=True))
INDIA_CALLING_CODE = '91'


def _phone_digits(value) -> tuple[str, bool]:
    if value is None or value == '':
        return '', False
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    elif isinstance(value, int):
        text = str(value)
    else:
        text = cell_text(value)
        if text.endswith('.0') and text[:-2].isdigit():
            text = text[:-2]
    stripped = text.strip()
    explicit = stripped.startswith('+') or stripped.startswith('00')
    digits = ''.join(ch for ch in stripped if ch.isdigit())
    if explicit and digits.startswith('00'):
        digits = digits[2:]
    return digits, explicit


def _split_calling_code(digits: str) -> tuple[str, str] | None:
    for code in _CALLING_CODES:
        if digits.startswith(code) and len(digits) - len(code) >= 6:
            return code, digits[len(code):]
    return None


def normalize_phone(value) -> str:
    """Store phones as '+<code> <national>'. Missing codes default to India (+91)."""
    digits, explicit = _phone_digits(value)
    if not digits:
        return ''
    code = ''
    national = digits
    if explicit:
        split = _split_calling_code(digits)
        if split:
            code, national = split
    elif len(digits) == 10:
        code, national = INDIA_CALLING_CODE, digits
    elif len(digits) == 12 and digits.startswith(INDIA_CALLING_CODE):
        code, national = INDIA_CALLING_CODE, digits[2:]
    else:
        split = _split_calling_code(digits)
        if split and len(digits) > 10:
            code, national = split
    if not code:
        code, national = INDIA_CALLING_CODE, digits
    return f'+{code} {national}'


def format_phone(value) -> str:
    return normalize_phone(value)


def clean_email(value) -> str:
    text = cell_text(value)
    if not text:
        return ''
    try:
        validate_email(text)
    except ValidationError:
        return ''
    return text


_TEXT_DATETIME_FORMATS = (
    '%d-%m-%Y',
    '%d/%m/%Y',
    '%d-%m-%Y %H:%M',
    '%d/%m/%Y %H:%M',
    '%d-%m-%Y %H:%M:%S',
    '%d/%m/%Y %H:%M:%S',
    '%Y-%m-%d %H:%M',
    '%Y-%m-%d %H:%M:%S',
)


def _parse_loose_datetime(value) -> datetime | None:
    text = str(value).strip()
    if not text:
        return None
    parsed = parse_datetime(text)
    if parsed is not None:
        return parsed
    only = parse_date(text)
    if only is not None:
        return datetime.combine(only, datetime.min.time())
    for fmt in _TEXT_DATETIME_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def as_date(value) -> date | None:
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = _parse_loose_datetime(value)
    return parsed.date() if parsed else None


def as_datetime(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime.min.time())
    else:
        parsed = _parse_loose_datetime(value)
        if parsed is None:
            return None
    if timezone.is_naive(parsed):
        return timezone.make_aware(parsed, timezone.get_current_timezone())
    return parsed


def find_sheet(workbook, *needles):
    lowered = [(name, name.lower()) for name in workbook.sheetnames]
    for needle in needles:
        want = needle.lower()
        for name, low in lowered:
            if want == low or want in low:
                return workbook[name]
    return None


def is_skip_row(values: list[str]) -> bool:
    if not any(values):
        return True
    if any(value.lower() in HEADER_MARKERS for value in values if value):
        return True
    if all(value in {'*', ''} for value in values):
        return True
    return False


def iter_data_rows(sheet: Worksheet):
    iterator = sheet.iter_rows(min_row=1, max_col=25, values_only=True)
    try:
        headers = header_index(next(iterator))
    except StopIteration:
        return
    for excel_row, raw in enumerate(iterator, start=2):
        yield excel_row, raw, headers


def parse_contact_row(raw, headers: dict[str, int]) -> dict | None:
    name = cell_text(cell_at(raw, headers, 'contact name', 'name'))
    company = cell_text(cell_at(raw, headers, 'company name', 'company'))
    if is_skip_row([name, company]):
        return None
    if not name and not company:
        return None
    return {
        'name': name[:255],
        'company_name': company[:255],
        'designation': cell_text(cell_at(raw, headers, 'designation'))[:150],
        'department': cell_text(cell_at(raw, headers, 'department'))[:150],
        'email': clean_email(cell_at(raw, headers, 'contact email', 'email')),
        'company_email': clean_email(cell_at(raw, headers, 'company email')),
        'phone': format_phone(cell_at(raw, headers, 'phone'))[:40],
        'address_1': cell_text(cell_at(raw, headers, 'address 1', 'address1', 'address'))[:255],
        'city': cell_text(cell_at(raw, headers, 'city', 'address city'))[:120],
        'state': cell_text(cell_at(raw, headers, 'state'))[:120],
        'country': cell_text(cell_at(raw, headers, 'country'))[:120],
        'gst_number': cell_text(cell_at(raw, headers, 'gst number', 'gst'))[:40],
        'created_at': as_datetime(cell_at(raw, headers, 'created at', 'created date')),
    }


def parse_opportunity_row(raw, headers: dict[str, int]) -> dict | None:
    name = cell_text(cell_at(raw, headers, 'prospect name', 'opportunity', 'name'))
    company = cell_text(cell_at(raw, headers, 'company name', 'company'))
    prospect = cell_text(cell_at(raw, headers, 'prospect number'))
    if is_skip_row([name, company, prospect]):
        return None
    expected, parsed_currency = parse_value(str(cell_at(raw, headers, 'expected amount') or ''))
    negotiated, _ = parse_value(str(cell_at(raw, headers, 'final negotiated amount') or ''))
    currency = coerce_currency(cell_text(cell_at(raw, headers, 'currency')) or parsed_currency or 'INR')
    return {
        'prospect_number': prospect[:32],
        'name': name[:255],
        'company_name': company[:255],
        'contact_name': cell_text(cell_at(raw, headers, 'contact names', 'contact name'))[:255],
        'lead_source': cell_text(cell_at(raw, headers, 'lead source'))[:150],
        'service_pillar': coerce_pillar(cell_text(cell_at(raw, headers, 'service pillar')))[:100],
        'sub_category': coerce_sub_category(
            cell_text(cell_at(raw, headers, 'sub-categories', 'sub categories', 'sub category'))
        )[:150],
        'delivery_partner': cell_text(cell_at(raw, headers, 'delivery partner'))[:150],
        'expected_amount': expected or Decimal('0'),
        'negotiated_amount': negotiated or Decimal('0'),
        'currency': currency,
        'stage': coerce_sales_stage(cell_text(cell_at(raw, headers, 'sales stage', 'stage'))),
        'owner_name': cell_text(cell_at(raw, headers, 'assigned to', 'owner')),
        'department': '',
        'follow_up_date': as_date(cell_at(raw, headers, 'followup date', 'follow up date')),
        'expected_close_date': as_date(cell_at(raw, headers, 'expected close date')),
        'comments': cell_text(cell_at(raw, headers, 'notes', 'comments')),
        'created_at': as_datetime(cell_at(raw, headers, 'created date', 'created at')),
    }


def contact_sheet(workbook) -> Worksheet | None:
    return find_sheet(workbook, 'contact')


def opportunity_sheet(workbook) -> Worksheet:
    found = find_sheet(workbook, 'lead - deal', 'lead', 'pipeline', 'opportunit')
    if found is not None:
        return found
    contacts = contact_sheet(workbook)
    for name in workbook.sheetnames:
        if contacts is None or name != contacts.title:
            return workbook[name]
    return workbook[workbook.sheetnames[0]]


def load_contacts_from_workbook(workbook) -> list[tuple[int, dict]]:
    sheet = contact_sheet(workbook) or workbook[workbook.sheetnames[0]]
    rows: list[tuple[int, dict]] = []
    for excel_row, raw, headers in iter_data_rows(sheet):
        parsed = parse_contact_row(raw, headers)
        if parsed:
            rows.append((excel_row, parsed))
    return rows


def load_opportunities_from_workbook(workbook) -> list[tuple[int, dict]]:
    sheet = opportunity_sheet(workbook)
    rows: list[tuple[int, dict]] = []
    for excel_row, raw, headers in iter_data_rows(sheet):
        parsed = parse_opportunity_row(raw, headers)
        if parsed:
            rows.append((excel_row, parsed))
    return rows


def _autosize(sheet: Worksheet, width: int = 22):
    for index in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _add_list_validation(sheet: Worksheet, col: int, options: list[str], rows: int = 200):
    formula = '"' + ','.join(options) + '"'
    if len(formula) > 250:
        return
    validation = DataValidation(type='list', formula1=formula, allow_blank=True)
    letter = get_column_letter(col)
    validation.add(f'{letter}2:{letter}{rows}')
    sheet.add_data_validation(validation)


def build_contact_workbook(rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Contact'
    sheet.append(CONTACT_HEADERS)
    for row in rows:
        sheet.append(row)
    _autosize(sheet)
    return _workbook_bytes(workbook)


def build_opportunity_workbook(rows: list[list], with_validation: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Lead - Deal'
    sheet.append(OPPORTUNITY_HEADERS)
    for row in rows:
        sheet.append(row)
    if with_validation:
        _add_list_validation(sheet, 6, SERVICE_PILLARS)
        _add_list_validation(sheet, 7, SUB_CATEGORIES)
        _add_list_validation(sheet, 10, CURRENCIES)
        _add_list_validation(sheet, 12, SALES_STAGES)
    _autosize(sheet)
    return _workbook_bytes(workbook)


def _workbook_bytes(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def open_upload(content: bytes):
    return load_workbook(io.BytesIO(content), data_only=True)


def contact_export_row(contact) -> list:
    return [
        contact.company.name,
        contact.name,
        contact.designation,
        contact.department,
        contact.email,
        contact.company_email,
        contact.phone,
        contact.address_1,
        contact.city,
        contact.state,
        contact.country,
        contact.gst_number,
        timezone.localtime(contact.created_at).strftime('%Y-%m-%d %H:%M'),
    ]


def opportunity_export_row(opportunity) -> list:
    return [
        opportunity.prospect_number,
        opportunity.name,
        opportunity.company.name,
        opportunity.contact.name if opportunity.contact_id else '',
        opportunity.lead_source,
        opportunity.service_pillar,
        opportunity.sub_category,
        opportunity.delivery_partner,
        float(opportunity.expected_amount),
        opportunity.currency or 'INR',
        float(opportunity.negotiated_amount) if opportunity.negotiated_amount else '',
        opportunity.stage,
        opportunity.owner_name,
        timezone.localtime(opportunity.created_at).strftime('%Y-%m-%d %H:%M'),
        timezone.localtime(opportunity.updated_at).strftime('%Y-%m-%d %H:%M'),
        opportunity.follow_up_date.isoformat() if opportunity.follow_up_date else '',
        opportunity.expected_close_date.isoformat() if opportunity.expected_close_date else '',
        opportunity.comments,
    ]


def dummy_contact_rows() -> list[list]:
    return [
        [
            'Chavans Technologies', 'Sumanth Chavan', 'CEO', 'Management',
            'sumanth.chavan@chavans.in', '', '919986047767', '',
            'Bengaluru', 'Karnataka', 'India', '', '2025-07-18 00:00',
        ],
        [
            'CitiusCloud Services', 'Roshan Shetty', '', 'Sales',
            'roshan.shetty@citiuscloud.com', '', '919800000000', '',
            'Mumbai', 'Maharashtra', 'India', '', '2025-08-22 00:00',
        ],
    ]


def dummy_opportunity_rows() -> list[list]:
    return [[
        'PRO-000',
        'Sample Cloud Migration',
        'Chavans Technologies',
        'Sumanth Chavan',
        'Referral',
        'Consulting',
        'Growth Advisory',
        '',
        600000,
        'INR',
        '',
        'Approach',
        'Avinash',
        '2025-08-22 00:00',
        '2025-08-22 00:00',
        '',
        '2025-09-19',
        'Sample row for the import template.',
    ]]
