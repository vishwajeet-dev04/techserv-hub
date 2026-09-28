import logging

from django.http import Http404, HttpResponse
from rest_framework import generics, status
from rest_framework.decorators import api_view, permission_classes, parser_classes
from rest_framework.parsers import JSONParser, MultiPartParser, FormParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from apps.pulse.constants import MAX_IMPORT_BYTES
from apps.pulse.models import Opportunity, OpportunityActivityAttachment, PulseContact, Reminder
from apps.pulse.serializers import (
    OpportunityActivitySerializer,
    OpportunitySerializer,
    OpsUpdateSerializer,
    PulseAssigneeSerializer,
    PulseContactSerializer,
    ReminderSerializer,
)
from apps.pulse.services import pulse_service
from common.permissions.roles import IsSalesOrAbove

logger = logging.getLogger(__name__)

XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
DOCX_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
MAX_PREVIEW_ROWS = 200
MAX_PREVIEW_COLS = 30
MAX_PREVIEW_TEXT_CHARS = 20000


def _collect_stage_files(request):
    files = []
    for key in ('files', 'attachments', 'file'):
        files.extend(request.FILES.getlist(key))
    # Deduplicate by object identity while preserving order
    seen = set()
    unique = []
    for item in files:
        marker = id(item)
        if marker in seen:
            continue
        seen.add(marker)
        unique.append(item)
    return unique


def _xlsx_response(content: bytes, filename: str) -> HttpResponse:
    response = HttpResponse(content, content_type=XLSX_TYPE)
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _read_upload(request):
    upload = request.FILES.get('file')
    if upload is None:
        return None, Response(
            {'error': True, 'message': 'Choose an Excel (.xlsx) file.', 'details': None},
            status=status.HTTP_400_BAD_REQUEST,
        )
    name = (upload.name or '').lower()
    if not name.endswith('.xlsx'):
        return None, Response(
            {'error': True, 'message': 'Only .xlsx files are accepted.', 'details': None},
            status=status.HTTP_400_BAD_REQUEST,
        )
    if upload.size and upload.size > MAX_IMPORT_BYTES:
        return None, Response(
            {'error': True, 'message': 'File is too large (max 2 MB).', 'details': None},
            status=status.HTTP_400_BAD_REQUEST,
        )
    content = upload.read()
    if len(content) > MAX_IMPORT_BYTES:
        return None, Response(
            {'error': True, 'message': 'File is too large (max 2 MB).', 'details': None},
            status=status.HTTP_400_BAD_REQUEST,
        )
    return content, None


class PulseContactListCreateView(generics.ListCreateAPIView):
    serializer_class = PulseContactSerializer
    permission_classes = [IsAuthenticated, IsSalesOrAbove]

    def get_queryset(self):
        return pulse_service.list_contacts(search=self.request.query_params.get('search'))

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        company_name = data.pop('company_name_input', '') or request.data.get('company_name', '')
        instance = pulse_service.create_contact(company_name=str(company_name), **data)
        logger.info('Pulse contact created: %s user=%s path=%s', instance.name, request.user.id, request.path)
        return Response(PulseContactSerializer(instance).data, status=status.HTTP_201_CREATED)


class PulseContactDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = PulseContactSerializer
    permission_classes = [IsAuthenticated, IsSalesOrAbove]

    def get_object(self):
        try:
            return pulse_service.get_contact(self.kwargs['pk'])
        except PulseContact.DoesNotExist as exc:
            raise Http404('Contact not found.') from exc

    def perform_update(self, serializer):
        data = serializer.validated_data
        company_name = data.pop('company_name_input', None)
        if company_name is None and 'company_name' in self.request.data:
            company_name = self.request.data.get('company_name')
        pulse_service.update_contact(serializer.instance, company_name=company_name, **data)

    def perform_destroy(self, instance):
        pulse_service.delete_contact(instance)


class OpportunityListCreateView(generics.ListCreateAPIView):
    serializer_class = OpportunitySerializer
    permission_classes = [IsAuthenticated, IsSalesOrAbove]

    def get_queryset(self):
        return pulse_service.list_opportunities(
            search=self.request.query_params.get('search'),
            stage=self.request.query_params.get('stage'),
        )

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        company_name = data.pop('company_name_input', '') or request.data.get('company_name', '')
        contact_name = data.pop('contact_name_input', '') or request.data.get('contact_name', '')
        owner_name = data.pop('owner_name', '') or request.data.get('owner_name', '')
        owner_id = request.data.get('owner')
        instance = pulse_service.create_opportunity(
            company_name=str(company_name),
            contact_name=str(contact_name),
            owner_name=str(owner_name or ''),
            owner_id=owner_id,
            **data,
        )
        logger.info('Pulse opportunity created: %s user=%s path=%s', instance.name, request.user.id, request.path)
        return Response(OpportunitySerializer(instance).data, status=status.HTTP_201_CREATED)


class OpportunityDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = OpportunitySerializer
    permission_classes = [IsAuthenticated, IsSalesOrAbove]
    parser_classes = [JSONParser, MultiPartParser, FormParser]

    def get_object(self):
        try:
            return pulse_service.get_opportunity(self.kwargs['pk'])
        except Opportunity.DoesNotExist as exc:
            raise Http404('Opportunity not found.') from exc

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context['request'] = self.request
        return context

    def perform_update(self, serializer):
        data = serializer.validated_data
        company_name = data.pop('company_name_input', None)
        contact_name = data.pop('contact_name_input', None)
        owner_name = data.pop('owner_name', None)
        stage_note = data.pop('stage_note', None)
        owner_id = self.request.data.get('owner')
        if company_name is None and 'company_name' in self.request.data:
            company_name = self.request.data.get('company_name')
        if contact_name is None and 'contact_name' in self.request.data:
            contact_name = self.request.data.get('contact_name')
        if stage_note is None and 'stage_note' in self.request.data:
            stage_note = self.request.data.get('stage_note')
        pulse_service.update_opportunity(
            serializer.instance,
            owner_name=owner_name,
            owner_id=owner_id,
            company_name=company_name,
            contact_name=contact_name,
            actor=self.request.user,
            stage_change_note=stage_note,
            attachment_files=_collect_stage_files(self.request),
            **data,
        )

    def perform_destroy(self, instance):
        pulse_service.delete_opportunity(instance)


class OpportunityActivityListView(generics.ListAPIView):
    serializer_class = OpportunityActivitySerializer
    permission_classes = [IsAuthenticated, IsSalesOrAbove]
    pagination_class = None

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context['request'] = self.request
        return context

    def get_queryset(self):
        return pulse_service.list_opportunity_activities(self.kwargs['pk'])


class OpsUpdateListView(generics.ListAPIView):
    serializer_class = OpsUpdateSerializer
    permission_classes = [IsAuthenticated, IsSalesOrAbove]
    pagination_class = None

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context['request'] = self.request
        return context

    def get_queryset(self):
        limit = self.request.query_params.get('limit') or 50
        try:
            limit = int(limit)
        except (TypeError, ValueError):
            limit = 50
        return pulse_service.list_ops_updates(limit=limit)


def _extract_xlsx_rows(file_field):
    from openpyxl import load_workbook

    file_field.open('rb')
    try:
        workbook = load_workbook(file_field, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        max_col = min(sheet.max_column or 1, MAX_PREVIEW_COLS)
        rows = []
        for row in sheet.iter_rows(max_row=MAX_PREVIEW_ROWS, max_col=max_col, values_only=True):
            rows.append(['' if cell is None else str(cell) for cell in row])
        return rows
    finally:
        file_field.close()


def _extract_docx_text(file_field):
    from docx import Document

    file_field.open('rb')
    try:
        document = Document(file_field)
        text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
        return text[:MAX_PREVIEW_TEXT_CHARS]
    finally:
        file_field.close()


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsSalesOrAbove])
def attachment_preview(request, pk):
    try:
        attachment = OpportunityActivityAttachment.objects.get(pk=pk)
    except OpportunityActivityAttachment.DoesNotExist as exc:
        raise Http404('Attachment not found.') from exc

    name = (attachment.file_name or '').lower()
    try:
        if name.endswith('.xlsx') or attachment.content_type == XLSX_TYPE:
            return Response({'kind': 'table', 'rows': _extract_xlsx_rows(attachment.file)})
        if name.endswith('.docx') or attachment.content_type == DOCX_TYPE:
            return Response({'kind': 'text', 'text': _extract_docx_text(attachment.file)})
    except Exception:
        logger.exception('Could not generate preview for attachment %s', pk)
        return Response(
            {'error': True, 'message': 'Could not generate a preview for this file.', 'details': None},
            status=status.HTTP_422_UNPROCESSABLE_ENTITY,
        )

    return Response(
        {'error': True, 'message': 'Preview not supported for this file type.', 'details': None},
        status=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
    )


class ReminderListCreateView(generics.ListCreateAPIView):
    serializer_class = ReminderSerializer
    permission_classes = [IsAuthenticated, IsSalesOrAbove]

    def get_queryset(self):
        return pulse_service.list_reminders(
            search=self.request.query_params.get('search'),
            status=self.request.query_params.get('status'),
            reminder_type=self.request.query_params.get('reminder_type'),
        )

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        instance = pulse_service.create_reminder(**serializer.validated_data)
        logger.info('Pulse reminder created: %s user=%s path=%s', instance.title, request.user.id, request.path)
        return Response(ReminderSerializer(instance).data, status=status.HTTP_201_CREATED)


class ReminderDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = ReminderSerializer
    permission_classes = [IsAuthenticated, IsSalesOrAbove]

    def get_object(self):
        try:
            return pulse_service.get_reminder(self.kwargs['pk'])
        except Reminder.DoesNotExist as exc:
            raise Http404('Reminder not found.') from exc

    def perform_update(self, serializer):
        pulse_service.update_reminder(serializer.instance, **serializer.validated_data)

    def perform_destroy(self, instance):
        pulse_service.delete_reminder(instance)


@api_view(['POST'])
@permission_classes([IsAuthenticated, IsSalesOrAbove])
def reminder_done(request, pk):
    try:
        reminder = pulse_service.get_reminder(pk)
    except Reminder.DoesNotExist:
        return Response(
            {'error': True, 'message': 'Reminder not found.', 'details': None},
            status=status.HTTP_404_NOT_FOUND,
        )
    instance = pulse_service.mark_reminder_done(reminder)
    return Response(ReminderSerializer(instance).data)


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsSalesOrAbove])
def pulse_assignees(request):
    return Response(PulseAssigneeSerializer(pulse_service.list_assignees(), many=True).data)


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsSalesOrAbove])
def pulse_summary(request):
    return Response(pulse_service.executive_summary())


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsSalesOrAbove])
def pulse_pipeline(request):
    return Response(
        pulse_service.pipeline_dashboard(
            date_from=request.query_params.get('from'),
            date_to=request.query_params.get('to'),
        )
    )


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsSalesOrAbove])
def pulse_executive(request):
    return Response(
        pulse_service.executive_dashboard(
            date_from=request.query_params.get('from'),
            date_to=request.query_params.get('to'),
        )
    )


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsSalesOrAbove])
def pulse_team(request):
    return Response(
        pulse_service.team_dashboard(
            date_from=request.query_params.get('from'),
            date_to=request.query_params.get('to'),
        )
    )


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsSalesOrAbove])
def pulse_settings(request):
    return Response(pulse_service.settings_payload())


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsSalesOrAbove])
def contacts_export(request):
    return _xlsx_response(pulse_service.export_contacts_xlsx(), 'pulse-contacts.xlsx')


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsSalesOrAbove])
def contacts_template(request):
    return _xlsx_response(pulse_service.contact_template_xlsx(), 'pulse-contacts-template.xlsx')


@api_view(['POST'])
@permission_classes([IsAuthenticated, IsSalesOrAbove])
@parser_classes([MultiPartParser, FormParser])
def contacts_import(request):
    content, error = _read_upload(request)
    if error is not None:
        return error
    result = pulse_service.import_contacts_xlsx(content)
    logger.info(
        'Pulse contacts imported created=%s errors=%s user=%s path=%s',
        result['created'], len(result['errors']), request.user.id, request.path,
    )
    return Response(result)


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsSalesOrAbove])
def opportunities_export(request):
    return _xlsx_response(pulse_service.export_opportunities_xlsx(), 'pulse-opportunities.xlsx')


@api_view(['GET'])
@permission_classes([IsAuthenticated, IsSalesOrAbove])
def opportunities_template(request):
    return _xlsx_response(pulse_service.opportunity_template_xlsx(), 'pulse-opportunities-template.xlsx')


@api_view(['POST'])
@permission_classes([IsAuthenticated, IsSalesOrAbove])
@parser_classes([MultiPartParser, FormParser])
def opportunities_import(request):
    content, error = _read_upload(request)
    if error is not None:
        return error
    result = pulse_service.import_opportunities_xlsx(content)
    logger.info(
        'Pulse opportunities imported created=%s errors=%s user=%s path=%s',
        result['created'], len(result['errors']), request.user.id, request.path,
    )
    return Response(result)
