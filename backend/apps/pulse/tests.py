from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from decimal import Decimal

from django.core.management import call_command
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from apps.pulse.constants import coerce_sales_stage
from apps.pulse.models import Opportunity, OpportunityActivity, PulseContact, ReminderStatus
from apps.users.models import User, UserRole


class PulseAPITests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            email='sales@test.com',
            name='Sales User',
            password='pass12345',
            role=UserRole.SALES,
        )
        self.avinash = User.objects.create_user(
            email='avinash@test.com',
            name='Avinash',
            password='pass12345',
            role=UserRole.SALES,
        )
        self.asif = User.objects.create_user(
            email='asif@test.com',
            name='Asif',
            password='pass12345',
            role=UserRole.SALES,
        )
        self.client.force_authenticate(self.user)

    def _contact_payload(self, **overrides):
        payload = {
            'name': 'Sumanth Chavan',
            'company_name': 'Chavans Technologies',
            'department': 'Management',
            'email': 'sumanth.chavan@chavans.in',
            'phone': '919986047767',
            'city': 'Bengaluru',
            'state': 'Karnataka',
            'country': 'India',
        }
        payload.update(overrides)
        return payload

    def _opportunity_payload(self, contact_id, **overrides):
        payload = {
            'name': 'CitiusCloud Services',
            'contact': contact_id,
            'owner_name': 'Avinash',
            'service_pillar': 'Consulting',
            'sub_category': 'Growth Advisory',
            'stage': 'Qualify',
            'expected_amount': '600000',
            'currency': 'INR',
            'expected_close_date': '2026-12-31',
        }
        payload.update(overrides)
        return payload

    def test_create_contact_requires_fields(self):
        response = self.client.post('/api/v1/pulse/contacts/', {'name': 'Only Name'}, format='json')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_create_contact_rejects_invalid_phone(self):
        response = self.client.post(
            '/api/v1/pulse/contacts/',
            {**self._contact_payload(), 'phone': '12345'},
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_create_contact_without_email_or_phone(self):
        response = self.client.post(
            '/api/v1/pulse/contacts/',
            self._contact_payload(email='', phone=''),
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(response.data['email'], '')
        self.assertEqual(response.data['phone'], '')

    def test_create_and_list_contact(self):
        response = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(), format='json')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['company_name'], 'Chavans Technologies')
        self.assertEqual(response.data['city'], 'Bengaluru')
        listed = self.client.get('/api/v1/pulse/contacts/')
        self.assertEqual(listed.status_code, status.HTTP_200_OK)
        self.assertEqual(listed.data['count'], 1)

    def test_create_opportunity_from_contact_dropdown(self):
        contact = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(), format='json')
        self.assertEqual(contact.status_code, status.HTTP_201_CREATED)
        response = self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(
                contact.data['id'],
                name='CitiusCloud Services',
                company_name='Chavans Technologies',
                contact_name='Sumanth Chavan',
                stage='Win',
            ),
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(response.data['prospect_number'], 'PRO-001')
        self.assertEqual(response.data['stage'], 'Win')
        self.assertEqual(response.data['status'], 'WON')
        self.assertEqual(response.data['currency'], 'INR')
        self.assertFalse(response.data['is_high_value'])
        second = self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(
                contact.data['id'],
                name='Second Deal',
                owner_name='Asif',
                service_pillar='Marketing',
                sub_category='ICP & ABM',
                stage='Approach or Demo',
                expected_amount='1700000',
                currency='USD',
            ),
            format='json',
        )
        self.assertEqual(second.status_code, status.HTTP_201_CREATED, second.data)
        self.assertEqual(second.data['prospect_number'], 'PRO-002')
        self.assertEqual(second.data['stage'], 'Approach')
        self.assertEqual(second.data['currency'], 'USD')
        summary = self.client.get('/api/v1/pulse/summary/')
        self.assertEqual(summary.data['opportunities'], 2)
        self.assertEqual(summary.data['won_opportunities'], 1)

    def test_opportunity_requires_currency_and_close_date(self):
        contact = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(), format='json')
        missing_currency = self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(contact.data['id'], currency=''),
            format='json',
        )
        self.assertEqual(missing_currency.status_code, status.HTTP_400_BAD_REQUEST)
        missing_close = self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(contact.data['id'], expected_close_date=''),
            format='json',
        )
        self.assertEqual(missing_close.status_code, status.HTTP_400_BAD_REQUEST)

    def test_opportunity_rejects_unknown_contact(self):
        response = self.client.post(
            '/api/v1/pulse/opportunities/',
            {
                'name': 'Ghost Deal',
                'company_name': 'Missing Co',
                'contact_name': 'Nobody',
                'owner_name': 'Avinash',
                'service_pillar': 'Consulting',
                'sub_category': 'Growth Advisory',
                'stage': 'Qualify',
                'expected_amount': '1000',
                'currency': 'INR',
                'expected_close_date': '2026-12-31',
            },
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_patch_contact_and_opportunity(self):
        contact = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(), format='json')
        patched = self.client.patch(
            f'/api/v1/pulse/contacts/{contact.data["id"]}/',
            {'department': 'Sales', 'company_name': 'Chavans Technologies', 'name': 'Sumanth Chavan'},
            format='json',
        )
        self.assertEqual(patched.status_code, status.HTTP_200_OK, patched.data)
        self.assertEqual(patched.data['department'], 'Sales')
        created = self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(contact.data['id'], stage='Prospect Qualification'),
            format='json',
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED, created.data)
        self.assertEqual(created.data['stage'], 'Qualify')
        owner_only = self.client.patch(
            f'/api/v1/pulse/opportunities/{created.data["id"]}/',
            {'owner_name': 'Asif'},
            format='json',
        )
        self.assertEqual(owner_only.status_code, status.HTTP_200_OK, owner_only.data)
        self.assertEqual(owner_only.data['owner_name'], 'Asif')
        blocked = self.client.patch(
            f'/api/v1/pulse/opportunities/{created.data["id"]}/',
            {'stage': 'Negotiation'},
            format='json',
        )
        self.assertEqual(blocked.status_code, status.HTTP_400_BAD_REQUEST)
        missing_amount = self.client.patch(
            f'/api/v1/pulse/opportunities/{created.data["id"]}/',
            {
                'stage': 'Negotiation',
                'follow_up_date': '2026-11-01',
                'stage_note': 'Moved into negotiation after demo.',
            },
            format='json',
        )
        self.assertEqual(missing_amount.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('negotiated_amount', missing_amount.data.get('details', missing_amount.data))
        updated = self.client.patch(
            f'/api/v1/pulse/opportunities/{created.data["id"]}/',
            {
                'stage': 'Negotiation',
                'follow_up_date': '2026-11-01',
                'stage_note': 'Moved into negotiation after demo.',
                'negotiated_amount': '250000',
            },
            format='json',
        )
        self.assertEqual(updated.status_code, status.HTTP_200_OK, updated.data)
        self.assertEqual(updated.data['stage'], 'Negotiation')
        self.assertEqual(Decimal(str(updated.data['negotiated_amount'])), Decimal('250000'))
        self.assertEqual(updated.data['prospect_number'], created.data['prospect_number'])
        activities = self.client.get(f'/api/v1/pulse/opportunities/{created.data["id"]}/activities/')
        self.assertEqual(activities.status_code, status.HTTP_200_OK)
        stage_rows = [row for row in activities.data if row['action'] == 'STAGE_CHANGE']
        self.assertEqual(len(stage_rows), 1)
        self.assertEqual(stage_rows[0]['from_stage'], 'Qualify')
        self.assertEqual(stage_rows[0]['to_stage'], 'Negotiation')

        fat_finger = self.client.patch(
            f'/api/v1/pulse/opportunities/{created.data["id"]}/',
            {
                'stage': 'Negotiation',
                'follow_up_date': '2026-11-01',
                'stage_note': 'Correcting negotiated amount.',
                'negotiated_amount': '11111111111',
            },
            format='json',
        )
        self.assertEqual(fat_finger.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('negotiated_amount', fat_finger.data.get('details', fat_finger.data))

    def test_stage_change_with_attachments(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from apps.pulse.models import OpportunityActivityAttachment

        contact = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(name='Attach Contact'), format='json')
        created = self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(contact.data['id'], name='Attach Deal', stage='Prospect'),
            format='json',
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED, created.data)
        file_a = SimpleUploadedFile('note.pdf', b'%PDF-1.4 demo', content_type='application/pdf')
        file_b = SimpleUploadedFile('shot.png', b'\x89PNG\r\n\x1a\n', content_type='image/png')
        updated = self.client.patch(
            f'/api/v1/pulse/opportunities/{created.data["id"]}/',
            {
                'stage': 'Demo',
                'follow_up_date': '2026-11-05',
                'stage_note': 'Shared proposal deck.',
                'files': [file_a, file_b],
            },
            format='multipart',
        )
        self.assertEqual(updated.status_code, status.HTTP_200_OK, updated.data)
        self.assertEqual(updated.data['stage'], 'Demo')
        activities = self.client.get(f'/api/v1/pulse/opportunities/{created.data["id"]}/activities/')
        self.assertEqual(activities.status_code, status.HTTP_200_OK)
        self.assertEqual(len(activities.data), 1)
        self.assertEqual(len(activities.data[0]['attachments']), 2)
        names = {row['file_name'] for row in activities.data[0]['attachments']}
        self.assertEqual(names, {'note.pdf', 'shot.png'})
        self.assertTrue(all(row.get('url') for row in activities.data[0]['attachments']))
        self.assertEqual(OpportunityActivityAttachment.objects.count(), 2)

        bad = self.client.patch(
            f'/api/v1/pulse/opportunities/{created.data["id"]}/',
            {
                'stage': 'Commercial',
                'follow_up_date': '2026-11-10',
                'stage_note': 'Bad file type.',
                'files': SimpleUploadedFile('virus.exe', b'MZ', content_type='application/octet-stream'),
            },
            format='multipart',
        )
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)

    def test_csv_attachment_preview_is_parsed_server_side(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from apps.pulse.models import OpportunityActivityAttachment

        contact = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(name='Csv Contact'), format='json')
        created = self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(contact.data['id'], name='Csv Deal', stage='Prospect'),
            format='json',
        )
        csv_bytes = '\ufeffName,Amount\n"Acme, Inc",1000\nBeta,250\n'.encode('utf-8')
        self.client.patch(
            f'/api/v1/pulse/opportunities/{created.data["id"]}/',
            {
                'stage': 'Demo',
                'follow_up_date': '2026-11-05',
                'stage_note': 'Shared pricing sheet.',
                'files': SimpleUploadedFile('pricing.csv', csv_bytes, content_type='text/csv'),
            },
            format='multipart',
        )
        attachment = OpportunityActivityAttachment.objects.get(file_name='pricing.csv')
        preview = self.client.get(f'/api/v1/pulse/attachments/{attachment.id}/preview/')
        self.assertEqual(preview.status_code, status.HTTP_200_OK, preview.data)
        self.assertEqual(preview.data['kind'], 'table')
        self.assertEqual(preview.data['rows'], [['Name', 'Amount'], ['Acme, Inc', '1000'], ['Beta', '250']])

    def test_cannot_delete_contact_with_opportunity(self):
        contact = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(), format='json')
        self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(contact.data['id'], name='Linked Deal', stage='Win'),
            format='json',
        )
        deleted = self.client.delete(f'/api/v1/pulse/contacts/{contact.data["id"]}/')
        self.assertEqual(deleted.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertTrue(deleted.data.get('error'))
        self.assertIn('opportunities', deleted.data.get('message', '').lower())
        self.assertTrue(PulseContact.objects.filter(pk=contact.data['id']).exists())

    def test_assignees_list_for_sales(self):
        response = self.client.get('/api/v1/pulse/assignees/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        names = {row['name'] for row in response.data}
        self.assertIn('Avinash', names)
        self.assertIn('Asif', names)

    def test_create_opportunity_with_owner_id(self):
        contact = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(), format='json')
        response = self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(
                contact.data['id'],
                name='Owner UUID Deal',
                owner_name='',
                owner=str(self.avinash.id),
                stage='Win',
            ),
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(response.data['owner_name'], 'Avinash')
        self.assertEqual(str(response.data['owner']), str(self.avinash.id))

    def test_create_opportunity_rejects_invalid_owner(self):
        contact = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(), format='json')
        response = self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(
                contact.data['id'],
                name='Bad Owner',
                owner_name='',
                owner='00000000-0000-0000-0000-000000000099',
                stage='Win',
            ),
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_reminder_requires_due_date(self):
        created = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(name='Priya Menon'), format='json')
        response = self.client.post(
            '/api/v1/pulse/reminders/',
            {
                'title': 'Follow up',
                'reminder_type': 'FOLLOW_UP',
                'contact': created.data['id'],
            },
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_reminder_mark_done(self):
        created = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(name='Priya Menon'), format='json')
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        reminder = self.client.post(
            '/api/v1/pulse/reminders/',
            {
                'title': 'Follow up',
                'reminder_type': 'FOLLOW_UP',
                'due_date': '2026-12-01',
                'contact': created.data['id'],
            },
            format='json',
        )
        self.assertEqual(reminder.status_code, status.HTTP_201_CREATED)
        done = self.client.post(f'/api/v1/pulse/reminders/{reminder.data["id"]}/done/')
        self.assertEqual(done.status_code, status.HTTP_200_OK)
        self.assertEqual(done.data['status'], ReminderStatus.DONE)

    def test_coerce_sales_stage_aliases(self):
        self.assertEqual(coerce_sales_stage('Prospect'), 'Prospect')
        self.assertEqual(coerce_sales_stage('Prospect Qualification'), 'Qualify')
        self.assertEqual(coerce_sales_stage('Approach or Demo'), 'Approach')
        self.assertEqual(coerce_sales_stage('Converted'), 'Commercial')
        self.assertEqual(coerce_sales_stage('Closed'), 'Win')

    def test_pipeline_dashboard_aggregates(self):
        contact = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(), format='json')
        self.assertEqual(contact.status_code, status.HTTP_201_CREATED)
        self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(
                contact.data['id'],
                name='Won Deal',
                stage='Win',
                lead_source='TechDay',
                expected_amount='1000000',
                negotiated_amount='900000',
            ),
            format='json',
        )
        self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(
                contact.data['id'],
                name='Lost Deal',
                owner_name='Asif',
                service_pillar='Marketing',
                sub_category='ICP & ABM',
                stage='Lost',
                lead_source='Sunfinity',
                expected_amount='500000',
            ),
            format='json',
        )
        self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(
                contact.data['id'],
                name='Open Deal',
                owner_name='Asif',
                service_pillar='Marketing',
                sub_category='ICP & ABM',
                stage='Approach',
                lead_source='Referral',
                expected_amount='400000',
            ),
            format='json',
        )
        response = self.client.get('/api/v1/pulse/pipeline/')
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data['contacts'], 1)
        self.assertEqual(response.data['opportunities'], 3)
        self.assertEqual(response.data['win_rate'], 50)
        self.assertEqual(response.data['pipeline_value'], 400000.0)
        self.assertEqual(response.data['open_pipeline']['count'], 1)
        self.assertEqual(response.data['open_pipeline']['value'], 400000.0)
        self.assertEqual(response.data['total_prospects'], 3)
        self.assertEqual(response.data['flow'][0]['label'], 'Total Opportunities')
        self.assertEqual(response.data['flow'][0]['key'], 'prospects')
        flow = {item['key']: item for item in response.data['flow']}
        self.assertEqual(flow['prospects']['count'], 3)
        self.assertEqual(flow['Win']['count'], 1)
        self.assertEqual(flow['Win']['value'], 900000.0)
        self.assertEqual(flow['Lost']['count'], 1)
        self.assertEqual(flow['Approach']['count'], 1)
        self.assertEqual(flow['Approach']['value'], 400000.0)
        sources = {item['name']: item['count'] for item in response.data['by_source']}
        self.assertEqual(sources['TechDay'], 1)
        self.assertEqual(sources['Sunfinity'], 1)
        bad = self.client.get('/api/v1/pulse/pipeline/', {'from': 'not-a-date'})
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)

    def test_team_dashboard_aggregates(self):
        contact = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(), format='json')
        self.assertEqual(contact.status_code, status.HTTP_201_CREATED)
        self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(
                contact.data['id'],
                name='Won Deal',
                stage='Win',
                expected_amount='1000000',
                negotiated_amount='900000',
            ),
            format='json',
        )
        self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(
                contact.data['id'],
                name='Open Lead',
                service_pillar='Marketing',
                sub_category='ICP & ABM',
                stage='Converted',
                expected_amount='200000',
            ),
            format='json',
        )
        self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(
                contact.data['id'],
                name='Lost Deal',
                owner_name='Asif',
                service_pillar='Marketing',
                sub_category='ICP & ABM',
                stage='Lost',
                expected_amount='500000',
            ),
            format='json',
        )
        response = self.client.get('/api/v1/pulse/team/')
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data['active_reps'], 2)
        self.assertEqual(response.data['total_leads'], 3)
        self.assertEqual(response.data['top_rep'], 'Avinash')
        self.assertEqual(response.data['total_won_value'], 900000.0)
        by_name = {row['name']: row for row in response.data['by_rep']}
        self.assertEqual(by_name['Avinash']['total_leads'], 2)
        self.assertEqual(by_name['Avinash']['converted'], 1)
        self.assertEqual(by_name['Avinash']['win_rate'], 100)
        self.assertEqual(by_name['Asif']['win_rate'], 0)
        bad = self.client.get('/api/v1/pulse/team/', {'from': 'not-a-date'})
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)

    def test_executive_dashboard(self):
        contact = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(), format='json')
        self.assertEqual(contact.status_code, status.HTTP_201_CREATED)
        contact_id = contact.data['id']
        self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(
                contact_id,
                name='Won Consulting',
                stage='Win',
                expected_amount='600000',
            ),
            format='json',
        )
        self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(
                contact_id,
                name='Closed Marketing',
                owner_name='Asif',
                service_pillar='Marketing',
                sub_category='ICP & ABM',
                stage='Closed',
                expected_amount='400000',
            ),
            format='json',
        )
        self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(
                contact_id,
                name='Lost Deal',
                owner_name='Asif',
                service_pillar='Marketing',
                sub_category='ICP & ABM',
                stage='Lost',
                expected_amount='100000',
            ),
            format='json',
        )
        self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(
                contact_id,
                name='Open Deal',
                stage='Approach or Demo',
                expected_amount='200000',
            ),
            format='json',
        )
        response = self.client.get('/api/v1/pulse/executive/')
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data['leads'], 1)
        self.assertEqual(response.data['deals'], 4)
        # Closed maps to Win; legacy Closed bucket is unused
        self.assertEqual(response.data['deals_closed'], 0)
        self.assertEqual(response.data['revenue_won'], 1000000.0)
        self.assertEqual(response.data['win_rate'], 67)
        self.assertIsNone(response.data['revenue_won_prior'])
        self.assertEqual(response.data['benchmarks']['avg_deal_size'], 500000.0)
        self.assertEqual(response.data['benchmarks']['lead_conversion'], 400)
        steps = {row['key']: row for row in response.data['funnel']['steps']}
        self.assertEqual(steps['win']['count'], 2)
        self.assertEqual(steps['closed']['count'], 0)
        self.assertIn('from_label', response.data['funnel']['drop'])
        self.assertIn('to_label', response.data['funnel']['drop'])
        self.assertIn('Consulting', {row['name'] for row in response.data['by_pillar_won']})
        self.assertTrue(response.data['volume'])
        far = self.client.get('/api/v1/pulse/executive/', {'from': '2099-01-01', 'to': '2099-01-31'})
        self.assertEqual(far.status_code, status.HTTP_200_OK, far.data)
        self.assertEqual(far.data['revenue_won'], 0.0)
        self.assertEqual(far.data['deals'], 0)
        self.assertIsNotNone(far.data['revenue_won_prior'])
        self.assertEqual(far.data['benchmarks']['avg_deal_size'], 500000.0)
        bad = self.client.get('/api/v1/pulse/executive/', {'from': 'not-a-date'})
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)

    def test_viewer_cannot_read(self):
        viewer = User.objects.create_user(
            email='view@test.com',
            name='Viewer',
            password='pass12345',
            role=UserRole.VIEWER,
        )
        self.client.force_authenticate(viewer)
        response = self.client.get('/api/v1/pulse/contacts/')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        pipeline = self.client.get('/api/v1/pulse/pipeline/')
        self.assertEqual(pipeline.status_code, status.HTTP_403_FORBIDDEN)
        team = self.client.get('/api/v1/pulse/team/')
        self.assertEqual(team.status_code, status.HTTP_403_FORBIDDEN)
        executive = self.client.get('/api/v1/pulse/executive/')
        self.assertEqual(executive.status_code, status.HTTP_403_FORBIDDEN)

    def test_contact_template_and_import(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        template = self.client.get('/api/v1/pulse/contacts/template/')
        self.assertEqual(template.status_code, status.HTTP_200_OK)
        self.assertIn('spreadsheet', template['Content-Type'])
        uploaded = SimpleUploadedFile(
            't.xlsx',
            template.content,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        imported = self.client.post('/api/v1/pulse/contacts/import/', {'file': uploaded}, format='multipart')
        self.assertEqual(imported.status_code, status.HTTP_200_OK, getattr(imported, 'data', imported.content))
        self.assertGreaterEqual(imported.data['created'], 1)
        self.assertEqual(PulseContact.objects.filter(name='Sumanth Chavan').count(), 1)
        exported = self.client.get('/api/v1/pulse/contacts/export/')
        book = load_workbook(BytesIO(exported.content))
        headers = [cell.value for cell in next(book.active.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(headers[0], 'Company Name')
        self.assertNotIn('ID', headers)
        self.assertNotIn('Contact ID', headers)

    def test_opportunity_import_creates_missing_contact(self):
        self.client.post('/api/v1/pulse/contacts/', self._contact_payload(), format='json')
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Lead - Deal'
        sheet.append([
            'Prospect Number', 'Prospect Name', 'Company name', 'Contact Names',
            'Lead Source', 'Service Pillar', 'Sub-Categories', 'Delivery Partner',
            'Expected Amount(₹)', 'Final Negotiated Amount(₹)', 'Sales Stage', 'Assigned To',
        ])
        sheet.append([
            'PRO-010', 'Sample Cloud Migration', 'Chavans Technologies', 'Sumanth Chavan',
            '', 'Consulting', 'Growth Advisory', '', 600000, '', 'Approach or Demo', 'Avinash',
        ])
        sheet.append([
            'PRO-011', 'Unknown Deal', 'Missing Co', 'Nobody',
            '', 'Consulting', 'Growth Advisory', '', 1000, '', 'Prospect Qualification', 'Avinash',
        ])
        buffer = BytesIO()
        workbook.save(buffer)
        from django.core.files.uploadedfile import SimpleUploadedFile
        uploaded = SimpleUploadedFile('opps.xlsx', buffer.getvalue())
        result = self.client.post('/api/v1/pulse/opportunities/import/', {'file': uploaded}, format='multipart')
        self.assertEqual(result.status_code, status.HTTP_200_OK, result.data)
        self.assertEqual(result.data['created'], 2)
        self.assertEqual(result.data['errors'], [])
        self.assertEqual(result.data['contacts_created'], 1)
        self.assertTrue(PulseContact.objects.filter(name='Nobody', company__name='Missing Co').exists())
        self.assertEqual(Opportunity.objects.get(prospect_number='PRO-010').owner, self.avinash)

    DEAL_HEADERS = [
        'Prospect Number', 'Prospect Name', 'Company name', 'Contact Names',
        'Lead Source', 'Service Pillar', 'Sub-Categories', 'Delivery Partner',
        'Expected Amount(₹)', 'Final Negotiated Amount(₹)', 'Sales Stage', 'Assigned To',
    ]

    def _pipeline_upload(self, deal_rows, contact_rows=None):
        from django.core.files.uploadedfile import SimpleUploadedFile

        workbook = Workbook()
        first = workbook.active
        if contact_rows is not None:
            first.title = 'Contact'
            first.append(['Company Name', 'Contact Name', 'Designation', 'Contact Email', 'Phone', 'City'])
            for row in contact_rows:
                first.append(row)
            deals = workbook.create_sheet('Lead - Deal')
        else:
            first.title = 'Lead - Deal'
            deals = first
        deals.append(self.DEAL_HEADERS)
        for row in deal_rows:
            deals.append(row)
        buffer = BytesIO()
        workbook.save(buffer)
        return self.client.post(
            '/api/v1/pulse/opportunities/import/',
            {'file': SimpleUploadedFile('pipeline.xlsx', buffer.getvalue())},
            format='multipart',
        )

    def test_opportunity_import_reads_contact_sheet_and_matches_loosely(self):
        self.client.post('/api/v1/pulse/contacts/', self._contact_payload(), format='json')
        result = self._pipeline_upload(
            deal_rows=[
                ['PRO-020', 'Loose Company', 'chavans technologies pvt. ltd.', 'sumanth  chavan',
                 '', 'Consulting', 'Growth Advisory', '', 1000, '', 'Prospect Qualification', 'asif'],
                ['PRO-021', 'Sheet Contact', 'Acme Corp', 'Ravi Kumar',
                 '', 'Consulting', 'Growth Advisory', '', 2000, '', 'Prospect Qualification', 'Ghost Person'],
                ['PRO-022', 'Email Owner', 'Acme Corp', 'Ravi Kumar',
                 '', 'Consulting', 'Growth Advisory', '', 3000, '', 'Prospect Qualification', 'AVINASH@test.com'],
            ],
            contact_rows=[['Acme', 'Ravi Kumar', 'CTO', 'ravi@acme.test', '9876543210', 'Pune']],
        )
        self.assertEqual(result.status_code, status.HTTP_200_OK, result.data)
        self.assertEqual(result.data['errors'], [])
        self.assertEqual(result.data['created'], 3)
        self.assertEqual(result.data['contacts_created'], 1)
        self.assertEqual(result.data['unmatched_assignees'], ['Ghost Person'])
        self.assertEqual(PulseContact.objects.filter(name__iexact='Sumanth Chavan').count(), 1)
        ravi = PulseContact.objects.get(name='Ravi Kumar')
        self.assertEqual(ravi.city, 'Pune')
        loose = Opportunity.objects.get(prospect_number='PRO-020')
        self.assertEqual(loose.owner, self.asif)
        self.assertEqual(loose.contact.name, 'Sumanth Chavan')
        ghost = Opportunity.objects.get(prospect_number='PRO-021')
        self.assertIsNone(ghost.owner)
        self.assertEqual(ghost.owner_name, 'Ghost Person')
        self.assertEqual(ghost.contact_id, ravi.id)
        self.assertEqual(Opportunity.objects.get(prospect_number='PRO-022').owner, self.avinash)
        self.assertTrue(any('Ghost Person' in w['message'] for w in result.data['warnings']))

    def test_opportunity_reimport_updates_by_prospect_number(self):
        row = ['PRO-030', 'Repeat Deal', 'Repeat Co', 'Rita', '', 'Consulting', 'Growth Advisory', '',
               1000, '', 'Prospect Qualification', 'Avinash']
        first = self._pipeline_upload([row])
        self.assertEqual(first.data['created'], 1, first.data)
        row[8] = 5000
        second = self._pipeline_upload([row])
        self.assertEqual(second.data['created'], 0, second.data)
        self.assertEqual(second.data['updated'], 1)
        self.assertEqual(second.data['contacts_created'], 0)
        self.assertEqual(Opportunity.objects.filter(prospect_number='PRO-030').count(), 1)
        self.assertEqual(Opportunity.objects.get(prospect_number='PRO-030').expected_amount, 5000)

    def test_opportunity_import_fills_blank_mandatory_fields(self):
        headers = self.DEAL_HEADERS + ['Expected Close Date']
        blank = ['PRO-050', '', '', '', '', '', '', '', None, '', '', '']
        real = [
            'PRO-051', 'Real Deal', 'Real Co', 'Ria', '', 'Consulting', 'Growth Advisory', '',
            2500, '', 'Prospect Qualification', 'Avinash', '',
        ]
        result = self._upload_deals(headers, [blank, real])
        self.assertEqual(result.status_code, status.HTTP_200_OK, result.data)
        self.assertEqual(result.data['errors'], [])
        self.assertEqual(result.data['created'], 2)
        self.assertTrue(result.data['warnings'])
        saved = Opportunity.objects.get(prospect_number='PRO-050')
        self.assertEqual(saved.expected_amount, 0)
        self.assertEqual(saved.service_pillar, 'NA')
        self.assertEqual(saved.sub_category, 'NA')
        self.assertIsNone(saved.expected_close_date)
        self.assertEqual(saved.name, 'NA')
        self.assertEqual(saved.company.name, 'NA')
        self.assertEqual(saved.contact.name, 'NA')
        self.assertEqual(saved.stage, 'Prospect')
        self.assertIsNone(saved.owner)
        self.assertEqual(saved.owner_name, 'NA')
        real_row = Opportunity.objects.get(prospect_number='PRO-051')
        self.assertEqual(real_row.expected_amount, 2500)
        kept = self.client.patch(
            f'/api/v1/pulse/opportunities/{saved.id}/',
            {'name': 'Filled later', 'service_pillar': 'NA', 'sub_category': 'NA'},
            format='json',
        )
        self.assertEqual(kept.status_code, status.HTTP_200_OK, kept.data)
        saved.refresh_from_db()
        self.assertEqual(saved.name, 'Filled later')
        self.assertEqual(saved.service_pillar, 'NA')
        rejected = self.client.patch(
            f'/api/v1/pulse/opportunities/{real_row.id}/',
            {'service_pillar': 'NA'},
            format='json',
        )
        self.assertEqual(rejected.status_code, status.HTTP_400_BAD_REQUEST)

    def test_opportunity_import_reports_every_error(self):
        bad = ['', 'No Pillar', 'Co', 'Person', '', 'Nonsense', 'Growth Advisory', '', 10, '', 'Prospect Qualification', 'Avinash']
        result = self._pipeline_upload([bad, list(bad), list(bad)])
        self.assertEqual(result.data['created'], 0)
        self.assertEqual(len(result.data['errors']), 3)
        self.assertEqual(result.data['errors'][0]['sheet'], 'Lead - Deal')
        self.assertFalse(PulseContact.objects.filter(name='Person').exists())

    def test_seed_pulse_loads_contact_sheet(self):
        workbook = Workbook()
        contacts = workbook.active
        contacts.title = 'Contact'
        contacts.append([
            'Company Name', 'Contact Name', 'Designation', 'Department',
            'Contact Email', 'Company Email', 'Phone', 'Address_1',
            'City', 'State', 'Country', 'GST Number',
        ])
        contacts.append(['*', '*', None, '*', '*', None, '*', None, '*', '*', '*', None])
        contacts.append([
            'Chavans Technologies', 'Sumanth Chavan', 'CEO', 'Management',
            'sumanth.chavan@chavans.in', None, 919986047767.0, None,
            'Bengaluru', 'Karnataka', 'India', None,
        ])
        deals = workbook.create_sheet('Lead - Deal')
        deals.append([
            'Prospect Number', 'Prospect Name', 'Company name ', 'Contact Names',
            'Lead Source', 'Service Pillar', 'Sub-Categories', 'Delivery Partner',
            'Expected Amount(₹)', 'Final Negotiated Amount(₹)', 'Sales Stage', 'Assigned To',
        ])
        deals.append(['*', '*', '*', '*', None, '*', '*', None, '*', None, None, '*'])
        deals.append([
            'PRO-001', 'CitiusCloud Services', 'Chavans Technologies', 'Sumanth Chavan',
            None, 'Consulting', 'Growth Advisory', None,
            '₹ 600,000.00', '600000.0', 'Win', 'Avinash',
        ])
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / 'pipeline.xlsx'
            workbook.save(path)
            call_command('seed_pulse', '--file', str(path), '--reset')
        contact = PulseContact.objects.get()
        self.assertEqual(contact.name, 'Sumanth Chavan')
        self.assertEqual(contact.phone, '+91 9986047767')
        self.assertEqual(contact.city, 'Bengaluru')
        opportunity = Opportunity.objects.get()
        self.assertEqual(opportunity.prospect_number, 'PRO-001')
        self.assertEqual(opportunity.name, 'CitiusCloud Services')
        self.assertEqual(opportunity.stage, 'Win')
        self.assertEqual(opportunity.contact_id, contact.id)
        self.assertEqual(opportunity.expected_amount, 600000)

    def _upload_deals(self, headers, rows):
        from django.core.files.uploadedfile import SimpleUploadedFile

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Lead - Deal'
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        buffer = BytesIO()
        workbook.save(buffer)
        return self.client.post(
            '/api/v1/pulse/opportunities/import/',
            {'file': SimpleUploadedFile('dated.xlsx', buffer.getvalue())},
            format='multipart',
        )

    def test_opportunity_import_uses_created_date_and_reimport_updates_it(self):
        from django.utils import timezone

        headers = self.DEAL_HEADERS + ['Created Date']
        row = [
            'PRO-040', 'Dated Deal', 'Dated Co', 'Dina', '', 'Consulting', 'Growth Advisory', '',
            1000, '', 'Prospect Qualification', 'Avinash', '22-08-2025',
        ]
        first = self._upload_deals(headers, [row])
        self.assertEqual(first.status_code, status.HTTP_200_OK, first.data)
        self.assertEqual(first.data['created'], 1)
        opportunity = Opportunity.objects.get(prospect_number='PRO-040')
        self.assertEqual(timezone.localtime(opportunity.created_at).date().isoformat(), '2025-08-22')
        self.assertFalse(
            OpportunityActivity.objects.filter(
                opportunity=opportunity, action='FIELD_UPDATE',
            ).exists()
        )

        row[-1] = '23-08-2025'
        second = self._upload_deals(headers, [row])
        self.assertEqual(second.data['updated'], 1, second.data)
        opportunity.refresh_from_db()
        self.assertEqual(timezone.localtime(opportunity.created_at).date().isoformat(), '2025-08-23')
        logged = OpportunityActivity.objects.get(opportunity=opportunity, action='FIELD_UPDATE')
        created_change = next(item for item in logged.changes if item['field'] == 'created_at')
        self.assertEqual(created_change['label'], 'Created date')
        self.assertTrue(created_change['new'].startswith('2025-08-23'))
        self.assertEqual(logged.actor_id, self.user.id)

        row[-1] = ''
        third = self._upload_deals(headers, [row])
        self.assertEqual(third.data['updated'], 1, third.data)
        opportunity.refresh_from_db()
        self.assertEqual(timezone.localtime(opportunity.created_at).date().isoformat(), '2025-08-23')
        self.assertEqual(
            OpportunityActivity.objects.filter(opportunity=opportunity, action='FIELD_UPDATE').count(),
            1,
        )

    def test_only_admin_can_change_created_date(self):
        from django.utils import timezone

        contact = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(), format='json')
        created = self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(contact.data['id']),
            format='json',
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED, created.data)
        url = f'/api/v1/pulse/opportunities/{created.data["id"]}/created-at/'
        denied = self.client.patch(url, {'created_at': '2025-08-22'}, format='json')
        self.assertEqual(denied.status_code, status.HTTP_403_FORBIDDEN)

        admin = User.objects.create_user(
            email='admin-dates@test.com',
            name='Date Admin',
            password='pass12345',
            role=UserRole.ADMIN,
        )
        self.client.force_authenticate(admin)
        updated = self.client.patch(url, {'created_at': '2025-08-22'}, format='json')
        self.assertEqual(updated.status_code, status.HTTP_200_OK, updated.data)
        opportunity = Opportunity.objects.get(pk=created.data['id'])
        self.assertEqual(timezone.localtime(opportunity.created_at).date().isoformat(), '2025-08-22')
        logged = OpportunityActivity.objects.get(opportunity=opportunity, action='FIELD_UPDATE')
        self.assertEqual(logged.actor_id, admin.id)
        self.assertEqual(logged.changes[0]['field'], 'created_at')
        self.assertTrue(logged.changes[0]['new'].startswith('2025-08-22'))

        same = self.client.patch(url, {'created_at': '2025-08-22'}, format='json')
        self.assertEqual(same.status_code, status.HTTP_200_OK, same.data)
        self.assertEqual(
            OpportunityActivity.objects.filter(opportunity=opportunity, action='FIELD_UPDATE').count(),
            1,
        )

    def test_opportunity_field_edit_is_logged_and_noop_is_not(self):
        contact = self.client.post('/api/v1/pulse/contacts/', self._contact_payload(name='Amount Contact'), format='json')
        created = self.client.post(
            '/api/v1/pulse/opportunities/',
            self._opportunity_payload(contact.data['id'], name='Amount Deal', expected_amount='1000'),
            format='json',
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED, created.data)
        url = f'/api/v1/pulse/opportunities/{created.data["id"]}/'
        changed = self.client.patch(url, {'expected_amount': '5000', 'owner_name': 'Asif'}, format='json')
        self.assertEqual(changed.status_code, status.HTTP_200_OK, changed.data)
        logged = OpportunityActivity.objects.get(opportunity_id=created.data['id'], action='FIELD_UPDATE')
        fields = {item['field']: item for item in logged.changes}
        self.assertEqual(fields['expected_amount']['old'], '1000.00')
        self.assertEqual(fields['expected_amount']['new'], '5000.00')
        self.assertEqual(fields['owner']['old'], 'Avinash')
        self.assertEqual(fields['owner']['new'], 'Asif')
        self.assertEqual(logged.actor_id, self.user.id)

        again = self.client.patch(url, {'expected_amount': '5000', 'name': 'Amount Deal'}, format='json')
        self.assertEqual(again.status_code, status.HTTP_200_OK, again.data)
        self.assertEqual(
            OpportunityActivity.objects.filter(opportunity_id=created.data['id']).count(),
            1,
        )
