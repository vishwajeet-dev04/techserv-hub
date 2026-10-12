"""Parse TechPay Sample MIS plate (Sales Register, Purchase Register, Bank book)."""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from apps.business.models import Company, CompanyStatus
from apps.finance.models import (
    CashflowAccount,
    CashflowBucket,
    CashflowDirection,
    CashflowEntry,
    Invoice,
    InvoiceStatus,
    Payable,
    PayableStatus,
    payable_due_date,
)

GST_MARKERS = ('gst', 'cgst', 'sgst', 'igst', 'unclaimed')
TDS_MARKERS = ('tds',)
SKIP_SALES_AMOUNT_COLS = {
    'date', 'particulars', 'buyer', 'buyer address', 'voucher type', 'voucher no.',
    'voucher ref. no.', 'gstin/uin', 'gross total',
}
SKIP_PURCHASE_META = {
    'date', 'particulars', 'supplier', 'supplier address', 'voucher type', 'voucher no.',
    'supplier invoice no.', 'supplier invoice date', 'gstin/uin', 'gross total',
}


def _norm(value) -> str:
    if value is None:
        return ''
    return str(value).strip()


def _norm_key(value) -> str:
    return _norm(value).lower().replace('\n', ' ')


def _as_decimal(value) -> Decimal:
    if value is None or value == '':
        return Decimal('0')
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).replace(',', '').replace('₹', '').strip()
    if not text or text in {'-', '—'}:
        return Decimal('0')
    try:
        return Decimal(text)
    except InvalidOperation:
        return Decimal('0')


def _as_date(value) -> date | None:
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _norm(value)
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%d-%b-%y', '%d-%b-%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _month_label(value: date | None) -> str:
    if not value:
        return ''
    return value.strftime('%B')


def _header_map(row: tuple) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(row):
        key = _norm_key(cell)
        if key and key not in mapping:
            mapping[key] = idx
    return mapping


def _find_header_row(rows: list[tuple], required: set[str]) -> tuple[int, dict[str, int]] | None:
    for index, row in enumerate(rows):
        mapping = _header_map(row)
        if required.issubset(mapping.keys()):
            return index, mapping
    return None


def _cell(row: tuple, headers: dict[str, int], *names: str):
    for name in names:
        idx = headers.get(name)
        if idx is not None and idx < len(row):
            return row[idx]
    return None


def _get_or_create_company(name: str, gstin: str = '') -> Company | None:
    cleaned = _norm(name)
    if not cleaned:
        return None
    company, _ = Company.objects.get_or_create(
        name=cleaned[:255],
        defaults={
            'status': CompanyStatus.CUSTOMER,
            'gst_number': _norm(gstin)[:40],
        },
    )
    if gstin and not company.gst_number:
        company.gst_number = _norm(gstin)[:40]
        company.save(update_fields=['gst_number', 'updated_at'])
    return company


def _load_rows(content: bytes) -> dict[str, list[tuple]]:
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    sheets: dict[str, list[tuple]] = {}
    for name in workbook.sheetnames:
        ws = workbook[name]
        sheets[name] = [tuple(row) for row in ws.iter_rows(values_only=True)]
    return sheets


def _find_sheet(sheets: dict[str, list[tuple]], *candidates: str) -> list[tuple] | None:
    lowered = {key.lower(): key for key in sheets}
    for candidate in candidates:
        hit = lowered.get(candidate.lower())
        if hit:
            return sheets[hit]
    return None


def _workspace_lookup(result: dict) -> dict:
    company, region = result.get('_workspace') or ('techserv', 'india')
    return {'workspace_company': company, 'workspace_region': region}


def _sales_invoice_number(row: tuple, headers: dict[str, int], invoice_date: date, buyer: str, gross: Decimal) -> str:
    voucher = _norm(_cell(row, headers, 'voucher no.'))
    if voucher:
        return voucher[:80]
    ref = _norm(_cell(row, headers, 'voucher ref. no.'))
    if ref:
        return ref[:80]
    slug = ''.join(ch for ch in buyer.upper() if ch.isalnum())[:24] or 'SALE'
    amount = str(gross).replace('.', '')
    return f'{invoice_date:%Y%m%d}-{slug}-{amount}'[:80]


def import_sales_register(rows: list[tuple], result: dict) -> None:
    found = _find_header_row(rows, {'date', 'gross total'})
    if not found or not any(key in found[1] for key in ('voucher no.', 'voucher ref. no.')):
        result['errors'].append({'sheet': 'Sales Register', 'message': 'Header row not found.'})
        return
    header_idx, headers = found
    amount_cols = []
    gst_cols = []
    for key, idx in headers.items():
        if key in SKIP_SALES_AMOUNT_COLS:
            continue
        if any(marker in key for marker in GST_MARKERS):
            gst_cols.append(idx)
        elif any(marker in key for marker in TDS_MARKERS):
            continue
        else:
            amount_cols.append(idx)

    last_date: date | None = None
    for excel_row, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        buyer = _norm(_cell(row, headers, 'buyer', 'particulars'))
        if not buyer or buyer.lower() in {'dr', 'cr'}:
            continue
        parsed_date = _as_date(_cell(row, headers, 'date'))
        if parsed_date:
            last_date = parsed_date
        invoice_date = parsed_date or last_date
        if not invoice_date:
            result['errors'].append({'sheet': 'Sales Register', 'row': excel_row, 'message': 'Missing date.'})
            continue
        gross = _as_decimal(_cell(row, headers, 'gross total'))
        if gross <= 0:
            continue
        voucher = _sales_invoice_number(row, headers, invoice_date, buyer, gross)
        gst = sum((_as_decimal(row[idx]) if idx < len(row) else Decimal('0')) for idx in gst_cols)
        taxable = sum((_as_decimal(row[idx]) if idx < len(row) else Decimal('0')) for idx in amount_cols)
        if taxable <= 0 and gross > 0:
            taxable = max(Decimal('0'), gross - gst)
        gstin = _norm(_cell(row, headers, 'gstin/uin'))
        company = _get_or_create_company(buyer, gstin)
        defaults = {
            'company': company,
            'customer_name': buyer[:255],
            'amount': gross,
            'taxable_amount': taxable,
            'gst_amount': gst,
            'invoice_date': invoice_date,
            'notes': f"MIS Particulars: {_norm(_cell(row, headers, 'particulars'))}".strip(),
        }
        obj, created = Invoice.objects.update_or_create(
            invoice_number=voucher[:80],
            **_workspace_lookup(result),
            defaults=defaults,
            create_defaults={**defaults, 'status': InvoiceStatus.DUE},
        )
        if created:
            result['invoices']['created'] += 1
        else:
            result['invoices']['updated'] += 1


def import_purchase_register(rows: list[tuple], result: dict) -> None:
    found = _find_header_row(rows, {'date', 'supplier', 'gross total'})
    if not found:
        result['errors'].append({'sheet': 'Purchase Register', 'message': 'Header row not found.'})
        return
    header_idx, headers = found
    expense_cols: list[tuple[str, int]] = []
    gst_cols: list[int] = []
    tds_cols: list[int] = []
    for key, idx in headers.items():
        if key in SKIP_PURCHASE_META:
            continue
        if any(marker in key for marker in TDS_MARKERS):
            tds_cols.append(idx)
        elif any(marker in key for marker in GST_MARKERS):
            gst_cols.append(idx)
        else:
            expense_cols.append((key, idx))

    last_date: date | None = None
    for excel_row, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        supplier = _norm(_cell(row, headers, 'supplier', 'particulars'))
        if not supplier or supplier.lower() in {'dr', 'cr', 'opening balance'}:
            continue
        invoice_no = _norm(_cell(row, headers, 'supplier invoice no.')) or _norm(_cell(row, headers, 'voucher no.'))
        if not invoice_no:
            continue
        supplier_invoice_date = _as_date(_cell(row, headers, 'supplier invoice date'))
        parsed_date = _as_date(_cell(row, headers, 'date'))
        if parsed_date:
            last_date = parsed_date
        entry_date = supplier_invoice_date or parsed_date or last_date
        gross = _as_decimal(_cell(row, headers, 'gross total'))
        gst = sum((_as_decimal(row[idx]) if idx < len(row) else Decimal('0')) for idx in gst_cols)
        tds = sum((_as_decimal(row[idx]) if idx < len(row) else Decimal('0')) for idx in tds_cols)
        taxable = Decimal('0')
        category = ''
        for name, idx in expense_cols:
            amount = _as_decimal(row[idx]) if idx < len(row) else Decimal('0')
            if amount > 0:
                taxable += amount
                if not category:
                    category = name.title()[:100]
        if taxable <= 0 and gross > 0:
            taxable = max(Decimal('0'), gross - gst)
        payable = max(Decimal('0'), gross - tds) if gross else taxable
        defaults = {
            'vendor': supplier[:255],
            'category': category or 'Purchase',
            'month_label': _month_label(entry_date),
            'taxable_amount': taxable,
            'gst_amount': gst,
            'tds_amount': tds,
            'total_amount': gross or (taxable + gst),
            'payable_amount': payable,
            'invoice_date': entry_date,
            'due_date': payable_due_date(entry_date),
            'notes': f"MIS Voucher: {_norm(_cell(row, headers, 'voucher no.'))}".strip(),
        }
        obj, created = Payable.objects.update_or_create(
            invoice_number=invoice_no[:80],
            vendor=supplier[:255],
            **_workspace_lookup(result),
            defaults=defaults,
            create_defaults={**defaults, 'status': PayableStatus.DUE},
        )
        if created:
            result['payables']['created'] += 1
        else:
            result['payables']['updated'] += 1


def import_bank_book(
    rows: list[tuple],
    result: dict,
    account: str = CashflowAccount.CURRENT,
    sheet_label: str = 'Bank book',
) -> None:
    header_idx = None
    headers = None
    for index, row in enumerate(rows):
        mapping = _header_map(row)
        if 'date' in mapping and ('vch type' in mapping or 'voucher type' in mapping):
            header_idx, headers = index, mapping
            break
    if header_idx is None or headers is None:
        result['errors'].append({'sheet': sheet_label, 'message': 'Header row not found.'})
        return

    vch_key = 'vch type' if 'vch type' in headers else 'voucher type'
    vch_no_key = 'vch no.' if 'vch no.' in headers else 'voucher no.'
    debit_key = 'debit' if 'debit' in headers else None
    credit_key = 'credit' if 'credit' in headers else None

    skip_party_idxs = {
        headers.get('date'),
        headers.get('category'),
        headers.get(vch_key),
        headers.get(vch_no_key),
        headers.get(debit_key) if debit_key else None,
        headers.get(credit_key) if credit_key else None,
    }
    skip_party_idxs.discard(None)

    last_date: date | None = None
    for excel_row, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        parsed_date = _as_date(_cell(row, headers, 'date'))
        if parsed_date:
            last_date = parsed_date

        party_bits = []
        particulars_cell = _norm(_cell(row, headers, 'particulars'))
        if particulars_cell and particulars_cell not in {'Dr', 'Cr', 'Educational'}:
            party_bits.append(particulars_cell)
        for idx in range(len(row)):
            if idx in skip_party_idxs:
                continue
            if idx == headers.get('particulars'):
                continue
            text = _norm(row[idx])
            if not text or text in {'Dr', 'Cr', 'Educational'}:
                continue
            if text.lower() in {b.lower() for b in party_bits}:
                continue
            party_bits.append(text)
        entity = next((bit for bit in party_bits if bit.lower() not in {'(as per details)'}), '')
        combined_particulars = ' '.join(party_bits).lower()
        if 'closing balance' in combined_particulars:
            continue

        is_opening = 'opening balance' in combined_particulars
        vch_type = _norm(_cell(row, headers, vch_key)).lower()
        if not vch_type and not is_opening:
            continue

        entry_date = parsed_date or last_date
        if not entry_date:
            result['errors'].append({'sheet': sheet_label, 'row': excel_row, 'message': 'Missing date.'})
            continue

        voucher_no = _norm(_cell(row, headers, vch_no_key))
        debit = _as_decimal(_cell(row, headers, debit_key) if debit_key else None)
        credit = _as_decimal(_cell(row, headers, credit_key) if credit_key else None)
        if debit <= 0 and credit <= 0:
            continue

        if is_opening:
            direction = CashflowDirection.INFLOW
            amount = debit if debit > 0 else credit
            bucket = CashflowBucket.CAPITAL
            vch_type = 'opening'
            category = 'Opening Balance'
            sub_category = 'Opening'
        elif vch_type == 'receipt':
            direction = CashflowDirection.INFLOW
            amount = debit if debit > 0 else credit
            bucket = CashflowBucket.REVENUE
            excel_category = _norm(_cell(row, headers, 'category'))
            category = excel_category or 'Others'
            sub_category = vch_type.title()
        elif vch_type == 'payment':
            direction = CashflowDirection.OUTFLOW
            amount = credit if credit > 0 else debit
            bucket = CashflowBucket.OPERATING
            excel_category = _norm(_cell(row, headers, 'category'))
            category = excel_category or 'Others'
            sub_category = vch_type.title()
        elif vch_type == 'contra':
            if debit > 0:
                direction = CashflowDirection.INFLOW
                amount = debit
            else:
                direction = CashflowDirection.OUTFLOW
                amount = credit
            bucket = CashflowBucket.CAPITAL
            excel_category = _norm(_cell(row, headers, 'category'))
            category = excel_category or 'Others'
            sub_category = vch_type.title()
        else:
            continue

        lookup = {
            'account': account,
            'entry_date': entry_date,
            'voucher_no': (voucher_no or ('OPENING' if is_opening else ''))[:80],
            'direction': direction,
            **_workspace_lookup(result),
        }
        defaults = {
            'bucket': bucket,
            'category': category[:100],
            'sub_category': sub_category[:100],
            'entity': entity[:255],
            'amount': amount,
            'notes': f'{sheet_label} {vch_type} {voucher_no}'.strip(),
        }
        existing = CashflowEntry.objects.filter(**lookup).first()
        if existing:
            for key, value in defaults.items():
                setattr(existing, key, value)
            existing.save()
            result['cashflow']['updated'] += 1
        else:
            CashflowEntry.objects.create(**lookup, **defaults)
            result['cashflow']['created'] += 1


def import_mis_workbook(content: bytes, workspace: tuple[str, str] | None = None) -> dict:
    result = {
        'invoices': {'created': 0, 'updated': 0},
        'payables': {'created': 0, 'updated': 0},
        'cashflow': {'created': 0, 'updated': 0},
        '_workspace': workspace or ('techserv', 'india'),
        'errors': [],
    }
    try:
        sheets = _load_rows(content)
    except Exception as exc:  # noqa: BLE001
        result['errors'].append({'sheet': '*', 'message': f'Could not open workbook: {exc}'})
        result.pop('_workspace')
        return result

    sales = _find_sheet(sheets, 'Sales Register', 'Sales Registry')
    purchase = _find_sheet(sheets, 'Purchase Register', 'Purchase Registry')
    bank = _find_sheet(sheets, 'Bank book', 'Bank Book', 'Operational Book')
    seed_bank = _find_sheet(
        sheets,
        'Seed Bank book',
        'Seed Bank Book',
        'Seed bank book',
        'Seed Book',
    )

    if sales is None:
        result['errors'].append({'sheet': 'Sales Register', 'message': 'Sheet missing.'})
    else:
        import_sales_register(sales, result)

    if purchase is None:
        result['errors'].append({'sheet': 'Purchase Register', 'message': 'Sheet missing.'})
    else:
        import_purchase_register(purchase, result)

    if bank is None:
        result['errors'].append({'sheet': 'Bank book', 'message': 'Sheet missing.'})
    else:
        import_bank_book(bank, result)

    if seed_bank is not None:
        import_bank_book(seed_bank, result, account=CashflowAccount.SEED, sheet_label='Seed Bank book')

    from apps.finance.status import refresh_auto_statuses

    refresh_auto_statuses(result.pop('_workspace'))
    return result


def import_mis_path(path: str | Path) -> dict:
    return import_mis_workbook(Path(path).read_bytes())
