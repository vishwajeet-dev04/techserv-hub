from datetime import date
from decimal import Decimal
from pathlib import Path

from rest_framework import status
from rest_framework.test import APITestCase

from apps.finance.models import FinancialMetric, FinancialPeriod, PeriodType
from apps.users.models import User, UserRole


class FinancialMetricAPITests(APITestCase):
    def setUp(self):
        self.finance_user = User.objects.create_user(
            email='finance@test.com',
            name='Finance User',
            password='pass12345',
            role=UserRole.FINANCE,
        )
        self.sales_user = User.objects.create_user(
            email='sales@test.com',
            name='Sales User',
            password='pass12345',
            role=UserRole.SALES,
        )
        self.client.force_authenticate(self.finance_user)
        self.period = FinancialPeriod.objects.create(
            period_type=PeriodType.MONTHLY,
            start_date=date(2026, 8, 1),
            end_date=date(2026, 8, 31),
            label='Aug 2026',
        )

    def test_create_manual_metric(self):
        response = self.client.post(
            '/api/v1/finance/metrics/',
            {
                'financial_period': str(self.period.id),
                'metric_code': 'revenue',
                'value': '5000000.50',
                'currency': 'INR',
            },
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['metric_code'], 'revenue')
        self.assertEqual(Decimal(response.data['value']), Decimal('5000000.50'))
        self.assertTrue(response.data['is_manual'])
        metric = FinancialMetric.objects.get(id=response.data['id'])
        self.assertEqual(metric.created_by, self.finance_user)
        self.assertTrue(metric.is_manual)

    def test_create_rejects_invalid_metric_code(self):
        response = self.client.post(
            '/api/v1/finance/metrics/',
            {
                'financial_period': str(self.period.id),
                'metric_code': 'not_a_real_metric',
                'value': '100',
                'currency': 'INR',
            },
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_create_rejects_non_numeric_value(self):
        response = self.client.post(
            '/api/v1/finance/metrics/',
            {
                'financial_period': str(self.period.id),
                'metric_code': 'revenue',
                'value': 'abc',
                'currency': 'INR',
            },
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_patch_updates_existing_metric_value(self):
        metric = FinancialMetric.objects.create(
            financial_period=self.period,
            metric_code='net_profit',
            value=Decimal('-1200.00'),
            currency='INR',
            is_manual=True,
            created_by=self.finance_user,
        )
        response = self.client.patch(
            f'/api/v1/finance/metrics/{metric.id}/',
            {'value': '2500.75', 'currency': 'INR'},
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(Decimal(response.data['value']), Decimal('2500.75'))
        self.assertTrue(response.data['is_manual'])
        metric.refresh_from_db()
        self.assertEqual(metric.value, Decimal('2500.75'))
        self.assertEqual(metric.metric_code, 'net_profit')

    def test_patch_rejects_non_numeric_value(self):
        metric = FinancialMetric.objects.create(
            financial_period=self.period,
            metric_code='cash',
            value=Decimal('10.00'),
            currency='INR',
            is_manual=True,
        )
        response = self.client.patch(
            f'/api/v1/finance/metrics/{metric.id}/',
            {'value': 'not-a-number'},
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        metric.refresh_from_db()
        self.assertEqual(metric.value, Decimal('10.00'))

    def test_sales_user_cannot_write_metrics(self):
        self.client.force_authenticate(self.sales_user)
        response = self.client.post(
            '/api/v1/finance/metrics/',
            {
                'financial_period': str(self.period.id),
                'metric_code': 'revenue',
                'value': '1',
                'currency': 'INR',
            },
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_sales_user_can_read_ops_summary(self):
        from apps.finance.models import PurchaseOrder, PurchaseOrderStatus
        PurchaseOrder.objects.create(
            po_number='PO-TEST-1',
            customer_name='Acme',
            total_value=Decimal('100000'),
            status=PurchaseOrderStatus.ACTIVE,
        )
        self.client.force_authenticate(self.sales_user)
        response = self.client.get('/api/v1/finance/ops/summary/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('po_metrics', response.data)
        self.assertEqual(response.data['po_metrics']['total_po_value'], 100000.0)

    def test_ops_summary_filters_by_date_range(self):
        from apps.finance.models import PurchaseOrder, PurchaseOrderStatus
        PurchaseOrder.objects.create(
            po_number='PO-IN-RANGE',
            customer_name='In range',
            total_value=Decimal('50000'),
            status=PurchaseOrderStatus.ACTIVE,
            po_date=date(2026, 4, 15),
        )
        PurchaseOrder.objects.create(
            po_number='PO-OUT-RANGE',
            customer_name='Out of range',
            total_value=Decimal('90000'),
            status=PurchaseOrderStatus.ACTIVE,
            po_date=date(2025, 1, 10),
        )
        self.client.force_authenticate(self.sales_user)
        response = self.client.get(
            '/api/v1/finance/ops/summary/',
            {'from': '2026-04-01', 'to': '2027-03-31'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['po_metrics']['total_po_value'], 50000.0)
        series = self.client.get(
            '/api/v1/finance/ops/dashboard-series/',
            {'from': '2026-04-01', 'to': '2026-06-30'},
        )
        self.assertEqual(series.status_code, status.HTTP_200_OK)
        months = [row['month'] for row in series.data.get('billing_vs_collections', [])]
        self.assertEqual(months, ['2026-04', '2026-05', '2026-06'])


class MisUploadAPITests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email='admin-mis@test.com',
            name='Admin',
            password='pass12345',
            role=UserRole.ADMIN,
        )
        self.finance_user = User.objects.create_user(
            email='finance-mis@test.com',
            name='Finance',
            password='pass12345',
            role=UserRole.FINANCE,
        )
        self.sales_user = User.objects.create_user(
            email='sales-mis@test.com',
            name='Sales',
            password='pass12345',
            role=UserRole.SALES,
        )
        self.client.force_authenticate(self.admin)

    def test_admin_uploads_xlsx(self):
        from io import BytesIO
        from pathlib import Path
        from tempfile import TemporaryDirectory

        from django.core.files.uploadedfile import SimpleUploadedFile
        from django.test import override_settings
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(['Sample'])
        buffer = BytesIO()
        workbook.save(buffer)
        uploaded = SimpleUploadedFile('sample-mis.xlsx', buffer.getvalue())
        with TemporaryDirectory() as tmp:
            with override_settings(MEDIA_ROOT=tmp):
                response = self.client.post('/api/v1/finance/mis/', {'file': uploaded}, format='multipart')
                self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
                self.assertEqual(response.data['file_name'], 'sample-mis.xlsx')
                self.assertEqual(response.data['file_type'], 'xlsx')
                self.assertGreater(response.data['byte_size'], 0)
                listed = self.client.get('/api/v1/finance/mis/')
                self.assertEqual(listed.data['count'], 1)
                stored = Path(tmp) / response.data['storage_path']
                self.assertTrue(stored.exists())

    def test_rejects_non_xlsx(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        uploaded = SimpleUploadedFile('notes.txt', b'hello')
        response = self.client.post('/api/v1/finance/mis/', {'file': uploaded}, format='multipart')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_finance_can_upload_and_sales_forbidden(self):
        from io import BytesIO
        from tempfile import TemporaryDirectory

        from django.core.files.uploadedfile import SimpleUploadedFile
        from django.test import override_settings
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(['Sample'])
        buffer = BytesIO()
        workbook.save(buffer)
        uploaded = SimpleUploadedFile('finance-mis.xlsx', buffer.getvalue())
        self.client.force_authenticate(self.finance_user)
        with TemporaryDirectory() as tmp:
            with override_settings(MEDIA_ROOT=tmp):
                response = self.client.post('/api/v1/finance/mis/', {'file': uploaded}, format='multipart')
                self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
                listed = self.client.get('/api/v1/finance/mis/')
                self.assertEqual(listed.status_code, status.HTTP_200_OK)
        self.client.force_authenticate(self.sales_user)
        self.assertEqual(self.client.get('/api/v1/finance/mis/').status_code, status.HTTP_403_FORBIDDEN)


SAMPLE_MIS = (
    Path(__file__).resolve().parents[3]
    / 'PROJECT_CONTEXT'
    / 'Files'
    / 'Sample data for MIS.xlsx'
)

SAMPLE_MIS_V2 = (
    Path(__file__).resolve().parents[3]
    / 'PROJECT_CONTEXT'
    / 'Files'
    / 'Sample data for MIS v2.xlsx'
)

DASHBOARD_SAMPLE_MIS = (
    Path(__file__).resolve().parents[3]
    / 'PROJECT_CONTEXT'
    / 'Files'
    / 'Dashboard Sample MIS.xlsx'
)


class MisImportParserTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email='admin-import@test.com',
            name='Admin',
            password='pass12345',
            role=UserRole.ADMIN,
        )
        self.finance_user = User.objects.create_user(
            email='finance-import@test.com',
            name='Finance',
            password='pass12345',
            role=UserRole.FINANCE,
        )

    def test_sample_mis_import_counts_and_idempotent(self):
        self.assertTrue(SAMPLE_MIS.exists(), f'Missing fixture: {SAMPLE_MIS}')
        from apps.finance.mis_import import import_mis_path
        from apps.finance.models import CashflowEntry, Invoice, Payable

        first = import_mis_path(SAMPLE_MIS)
        self.assertEqual(first['errors'], [])
        self.assertEqual(first['invoices']['created'], 7)
        self.assertEqual(first['payables']['created'], 11)
        self.assertEqual(first['cashflow']['created'], 22)
        self.assertEqual(Invoice.objects.count(), 7)
        self.assertEqual(Payable.objects.count(), 11)
        self.assertEqual(CashflowEntry.objects.count(), 22)
        self.assertTrue(
            Payable.objects.filter(tds_amount__gt=0).exists(),
            'Expected at least one payable with TDS from Purchase Register',
        )

        second = import_mis_path(SAMPLE_MIS)
        self.assertEqual(second['errors'], [])
        self.assertEqual(second['invoices']['created'], 0)
        self.assertEqual(second['invoices']['updated'], 7)
        self.assertEqual(second['payables']['updated'], 11)
        self.assertEqual(second['cashflow']['updated'], 22)
        self.assertEqual(Invoice.objects.count(), 7)
        self.assertEqual(Payable.objects.count(), 11)
        self.assertEqual(CashflowEntry.objects.count(), 22)

    def test_import_sets_open_status_and_keeps_manual_status(self):
        from apps.finance.mis_import import import_mis_path
        from apps.finance.models import (
            OPEN_STATUSES,
            FinanceStatusChange,
            Invoice,
            InvoiceStatus,
            Payable,
            PayableStatus,
        )
        from apps.finance.status import change_status

        import_mis_path(SAMPLE_MIS)

        self.assertFalse(Payable.objects.exclude(status__in=OPEN_STATUSES).exists())
        self.assertFalse(Invoice.objects.exclude(status__in=OPEN_STATUSES).exists())
        tarun = Payable.objects.get(vendor='Tarun P S')
        self.assertEqual(tarun.total_amount, Decimal('55037.25'))

        change_status(tarun, PayableStatus.PAID, self.finance_user, 'Paid by NEFT')
        savex = Invoice.objects.filter(customer_name__icontains='Savex').first()
        self.assertIsNotNone(savex)
        change_status(savex, InvoiceStatus.RECEIVED, self.finance_user)
        savex.refresh_from_db()
        self.assertEqual(savex.amount_received, savex.amount)

        import_mis_path(SAMPLE_MIS)
        tarun.refresh_from_db()
        savex.refresh_from_db()
        self.assertEqual(tarun.status, PayableStatus.PAID)
        self.assertTrue(tarun.status_manual)
        self.assertEqual(savex.status, InvoiceStatus.RECEIVED)
        self.assertEqual(FinanceStatusChange.objects.count(), 2)

    def test_bad_sheet_reports_errors(self):
        from io import BytesIO

        from apps.finance.mis_import import import_mis_workbook
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.title = 'Not A Ledger'
        workbook.active.append(['Nope'])
        buffer = BytesIO()
        workbook.save(buffer)
        result = import_mis_workbook(buffer.getvalue())
        messages = {row['sheet']: row['message'] for row in result['errors']}
        self.assertIn('Sales Register', messages)
        self.assertIn('Purchase Register', messages)
        self.assertIn('Bank book', messages)

    def test_upload_parses_into_ledgers_and_summary(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from django.test import override_settings
        from tempfile import TemporaryDirectory

        from apps.finance.models import CashflowEntry, Invoice, Payable

        self.client.force_authenticate(self.admin)
        uploaded = SimpleUploadedFile(
            'Sample data for MIS.xlsx',
            SAMPLE_MIS.read_bytes(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        with TemporaryDirectory() as tmp:
            with override_settings(MEDIA_ROOT=tmp):
                response = self.client.post('/api/v1/finance/mis/', {'file': uploaded}, format='multipart')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        result = response.data.get('import_result') or {}
        self.assertEqual(result.get('errors'), [])
        self.assertEqual(Invoice.objects.count(), 7)
        self.assertEqual(Payable.objects.count(), 11)
        self.assertEqual(CashflowEntry.objects.count(), 22)

        self.client.force_authenticate(self.finance_user)
        cashflow = self.client.get('/api/v1/finance/ledger/summary/', {'page': 'cashflow'})
        self.assertEqual(cashflow.status_code, status.HTTP_200_OK)
        self.assertGreater(cashflow.data['total_inflow'], 0)
        payables = self.client.get('/api/v1/finance/ledger/summary/', {'page': 'payables'})
        self.assertEqual(payables.status_code, status.HTTP_200_OK)
        self.assertGreater(payables.data['total_payable'], 0)
        invoices = self.client.get('/api/v1/finance/ledger/summary/', {'page': 'invoices'})
        self.assertEqual(invoices.status_code, status.HTTP_200_OK)
        self.assertEqual(invoices.data['count'], 7)
        filtered = self.client.get('/api/v1/finance/cashflow/', {'direction': 'INFLOW'})
        self.assertEqual(filtered.status_code, status.HTTP_200_OK)
        self.assertTrue(all(row['direction'] == 'INFLOW' for row in filtered.data['results']))

    def test_sample_mis_v2_imports_bank_category_column(self):
        self.assertTrue(SAMPLE_MIS_V2.exists(), f'Missing fixture: {SAMPLE_MIS_V2}')
        from apps.finance.mis_import import import_mis_path
        from apps.finance.models import CashflowEntry

        result = import_mis_path(SAMPLE_MIS_V2)
        self.assertEqual(result['errors'], [])
        self.assertGreater(result['cashflow']['created'], 0)

        categories = set(CashflowEntry.objects.values_list('category', flat=True))
        self.assertIn('Salaries', categories)
        self.assertTrue(
            any(name.lower().startswith('customer receipt') for name in categories),
            f'Expected Customer receipt category, got {categories}',
        )
        # Category should come from Excel, not the party name for categorized rows
        salary_row = CashflowEntry.objects.filter(category='Salaries').first()
        self.assertIsNotNone(salary_row)
        self.assertNotEqual(salary_row.category, salary_row.entity)
        self.assertTrue(
            salary_row.entity and 'praveen' in salary_row.entity.lower(),
            f'Expected Particulars/entity to be salary party, got {salary_row.entity!r}',
        )

        fee_row = CashflowEntry.objects.filter(
            category__iexact='Professional fee',
            entity__icontains='Tarun',
        ).first()
        self.assertIsNotNone(fee_row)
        self.assertEqual(fee_row.category, 'Professional fee')
        self.assertIn('Tarun', fee_row.entity)
        self.assertNotEqual(fee_row.category.lower(), fee_row.entity.lower())

        self.client.force_authenticate(self.finance_user)
        summary = self.client.get('/api/v1/finance/ledger/summary/', {'page': 'cashflow'})
        self.assertEqual(summary.status_code, status.HTTP_200_OK)
        inflow_cards = summary.data['inflow_cards']
        outflow_cards = summary.data['outflow_cards']
        card_names = {row['category'] for row in inflow_cards + outflow_cards}
        self.assertIn('Salaries', {row['category'] for row in outflow_cards})
        self.assertNotIn('TOTAL (Net)', card_names)
        self.assertNotIn('Opening Balance', card_names)
        self.assertTrue(all(row['direction'] == 'Inflow' for row in inflow_cards))
        self.assertTrue(all(row['direction'] == 'Outflow' for row in outflow_cards))
        for cards in (inflow_cards, outflow_cards):
            amounts = [row['amount'] for row in cards]
            self.assertEqual(amounts, sorted(amounts, reverse=True))
        # Party names must not appear as category card labels
        self.assertFalse(any('tarun' in name.lower() for name in card_names))
        self.assertFalse(any('savex' in name.lower() for name in card_names))

    def test_opening_balance_imported_and_bank_position(self):
        from apps.finance.mis_import import import_mis_path
        from apps.finance.models import CashflowEntry

        fixture = SAMPLE_MIS_V2 if SAMPLE_MIS_V2.exists() else SAMPLE_MIS
        self.assertTrue(fixture.exists(), f'Missing fixture: {fixture}')
        result = import_mis_path(fixture)
        self.assertEqual(result['errors'], [])

        opening = CashflowEntry.objects.filter(sub_category__iexact='Opening').first()
        self.assertIsNotNone(opening)
        self.assertEqual(opening.category, 'Opening Balance')
        self.assertEqual(opening.direction, 'INFLOW')
        self.assertGreater(opening.amount, 0)

        self.client.force_authenticate(self.finance_user)
        response = self.client.get('/api/v1/finance/bank-position/')
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        ops = response.data['operational']
        self.assertTrue(ops['available'])
        self.assertAlmostEqual(ops['opening'], float(opening.amount), places=2)
        self.assertGreater(ops['revenue_receipts'], 0)
        self.assertGreater(ops['all_outflows'], 0)
        expected_closing = (
            ops['opening']
            + ops['revenue_receipts']
            + ops['transfer_in']
            - ops['all_outflows']
            - ops['transfer_out']
        )
        self.assertAlmostEqual(ops['closing'], expected_closing, places=2)
        self.assertAlmostEqual(response.data['combined_closing'], ops['closing'], places=2)
        self.assertFalse(response.data['seed']['available'])
        self.assertEqual(response.data['seed']['status'], 'awaiting_seed_bank_book')
        if ops['transfer_in'] > 0:
            self.assertAlmostEqual(
                response.data['seed']['transfer_to_ops'],
                ops['transfer_in'],
                places=2,
            )
            self.assertAlmostEqual(response.data['internal_transfer'], ops['transfer_in'], places=2)

    def test_seed_bank_book_sheet_imports_into_seed_account(self):
        from io import BytesIO

        from openpyxl import load_workbook

        from apps.finance.mis_import import import_mis_workbook
        from apps.finance.models import CashflowAccount, CashflowEntry

        workbook = load_workbook(SAMPLE_MIS)
        bank_name = next(name for name in workbook.sheetnames if name.lower() == 'bank book')
        seed_sheet = workbook.copy_worksheet(workbook[bank_name])
        seed_sheet.title = 'Seed Bank book'
        buffer = BytesIO()
        workbook.save(buffer)

        result = import_mis_workbook(buffer.getvalue())
        self.assertEqual(result['errors'], [])
        current = CashflowEntry.objects.filter(account=CashflowAccount.CURRENT).count()
        seed = CashflowEntry.objects.filter(account=CashflowAccount.SEED).count()
        self.assertEqual(current, 22)
        self.assertEqual(seed, 22)

        again = import_mis_workbook(buffer.getvalue())
        self.assertEqual(again['cashflow']['created'], 0)
        self.assertEqual(CashflowEntry.objects.count(), 44)

        self.client.force_authenticate(self.finance_user)
        seed_list = self.client.get('/api/v1/finance/cashflow/', {'account': 'SEED', 'page_size': 5})
        self.assertEqual(seed_list.status_code, status.HTTP_200_OK)
        self.assertEqual(seed_list.data['count'], 22)
        self.assertEqual(len(seed_list.data['results']), 5)
        self.assertTrue(all(row['account'] == 'SEED' for row in seed_list.data['results']))
        default_list = self.client.get('/api/v1/finance/cashflow/')
        self.assertTrue(all(row['account'] == 'CURRENT' for row in default_list.data['results']))

        seed_summary = self.client.get('/api/v1/finance/ledger/summary/', {'page': 'cashflow', 'account': 'SEED'})
        self.assertTrue(seed_summary.data['has_entries'])
        self.assertEqual(seed_summary.data['account'], 'SEED')
        position = self.client.get('/api/v1/finance/bank-position/', {'account': 'SEED'})
        self.assertTrue(position.data['seed']['available'])
        self.assertEqual(position.data['seed']['status'], 'available')
        self.assertEqual(position.data['combined_scope'], 'seed')

    def test_seed_account_empty_without_sheet(self):
        from apps.finance.mis_import import import_mis_path

        import_mis_path(SAMPLE_MIS)
        self.client.force_authenticate(self.finance_user)
        summary = self.client.get('/api/v1/finance/ledger/summary/', {'page': 'cashflow', 'account': 'SEED'})
        self.assertFalse(summary.data['has_entries'])
        self.assertEqual(summary.data['inflow_cards'], [])
        self.assertEqual(summary.data['outflow_cards'], [])

    def test_dashboard_sample_mis_operational_and_seed_books(self):
        self.assertTrue(DASHBOARD_SAMPLE_MIS.exists(), f'Missing fixture: {DASHBOARD_SAMPLE_MIS}')
        from apps.finance.mis_import import import_mis_path
        from apps.finance.models import CashflowAccount, CashflowEntry, Invoice, Payable

        first = import_mis_path(DASHBOARD_SAMPLE_MIS)
        self.assertGreater(first['invoices']['created'], 0)
        self.assertGreater(first['payables']['created'], 0)
        self.assertGreater(first['cashflow']['created'], 0)
        self.assertGreater(Invoice.objects.count(), 0)
        self.assertGreater(Payable.objects.count(), 0)
        self.assertTrue(CashflowEntry.objects.filter(account=CashflowAccount.CURRENT).exists())
        self.assertTrue(CashflowEntry.objects.filter(account=CashflowAccount.SEED).exists())
        self.assertTrue(Invoice.objects.exclude(invoice_number='').exists())

        invoices = Invoice.objects.count()
        payables = Payable.objects.count()
        cashflow = CashflowEntry.objects.count()
        second = import_mis_path(DASHBOARD_SAMPLE_MIS)
        self.assertEqual(second['invoices']['created'], 0)
        self.assertEqual(second['invoices']['updated'], invoices)
        self.assertEqual(second['payables']['created'], 0)
        self.assertEqual(Invoice.objects.count(), invoices)
        self.assertEqual(Payable.objects.count(), payables)
        self.assertEqual(CashflowEntry.objects.count(), cashflow)


class MasterDataAPITests(APITestCase):
    def setUp(self):
        from apps.finance.models import FinanceClient, FinanceVendor

        FinanceClient.objects.all().delete()
        FinanceVendor.objects.all().delete()
        self.finance_user = User.objects.create_user(
            email='md-finance@test.com',
            name='MD Finance',
            password='pass12345',
            role=UserRole.FINANCE,
        )
        self.viewer = User.objects.create_user(
            email='md-viewer@test.com',
            name='MD Viewer',
            password='pass12345',
            role=UserRole.VIEWER,
        )
        self.client.force_authenticate(self.finance_user)

    def _upload(self, url, content, name='data.xlsx'):
        from django.core.files.uploadedfile import SimpleUploadedFile

        upload = SimpleUploadedFile(
            name,
            content,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        return self.client.post(url, {'file': upload}, format='multipart')

    def test_client_crud_and_pan_autofill(self):
        response = self.client.post(
            '/api/v1/finance/clients/',
            {
                'name': '  Acme   Corp ',
                'service_line': 'Managed Services',
                'email': 'ap@acme.in',
                'mobile': '+91 98765 43210',
                'gstin': '29abcde1234f1z5',
                'msme_registered': True,
                'msme_number': 'udyam-kr-01-0000001',
            },
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(response.data['name'], 'Acme Corp')
        self.assertEqual(response.data['gstin'], '29ABCDE1234F1Z5')
        self.assertEqual(response.data['pan'], 'ABCDE1234F')
        self.assertEqual(response.data['mobile'], '+919876543210')
        client_id = response.data['id']

        response = self.client.patch(
            f'/api/v1/finance/clients/{client_id}/', {'is_active': False}, format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertFalse(response.data['is_active'])

        response = self.client.get('/api/v1/finance/clients/', {'active': 'true'})
        self.assertEqual(response.data['count'], 0)
        response = self.client.get('/api/v1/finance/clients/', {'search': 'acme'})
        self.assertEqual(response.data['count'], 1)

        response = self.client.delete(f'/api/v1/finance/clients/{client_id}/')
        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)

    def test_client_validation(self):
        self.client.post('/api/v1/finance/clients/', {'name': 'Acme'}, format='json')
        response = self.client.post(
            '/api/v1/finance/clients/',
            {'name': 'ACME', 'pan': 'BAD', 'gstin': '123', 'mobile': '12345', 'email': 'nope'},
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        details = response.data.get('details') or response.data
        for field in ('name', 'pan', 'gstin', 'mobile', 'email'):
            self.assertIn(field, details)

    def test_vendor_tds_section_required(self):
        response = self.client.post(
            '/api/v1/finance/vendors/',
            {'name': 'Landlord LLP', 'tds_applicable': True},
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        details = response.data.get('details') or response.data
        self.assertIn('tds_section', details)

        response = self.client.post(
            '/api/v1/finance/vendors/',
            {'name': 'Landlord LLP', 'tds_applicable': True, 'tds_section': '194i', 'tds_rate': '10'},
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(response.data['tds_section'], '194I')

    def test_list_rows_include_ledger_stats(self):
        from apps.finance.models import FinanceClient, FinanceVendor, Invoice, Payable, PayableStatus

        FinanceClient.objects.create(name='Loose Customer')
        Invoice.objects.create(
            invoice_number='INV-1', customer_name='loose customer', amount=Decimal('1000'),
            amount_received=Decimal('400'), invoice_date=date(2026, 8, 1),
        )
        FinanceVendor.objects.create(name='Rent Co')
        Payable.objects.create(vendor='Rent Co', total_amount=500, payable_amount=500)
        Payable.objects.create(
            vendor='rent co', total_amount=300, payable_amount=300, status=PayableStatus.PAID,
        )

        row = self.client.get('/api/v1/finance/clients/').data['results'][0]
        self.assertEqual(row['total_invoiced'], 1000.0)
        self.assertEqual(row['total_paid'], 400.0)
        row = self.client.get('/api/v1/finance/vendors/').data['results'][0]
        self.assertEqual(row['total_payable'], 800.0)
        self.assertEqual(row['outstanding'], 500.0)

    def test_viewer_is_read_only(self):
        self.client.force_authenticate(self.viewer)
        response = self.client.get('/api/v1/finance/vendors/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        response = self.client.post('/api/v1/finance/vendors/', {'name': 'X'}, format='json')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        from apps.finance.master_excel import vendor_template_xlsx

        response = self._upload('/api/v1/finance/vendors/import/', vendor_template_xlsx())
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_vendor_import_updates_by_name(self):
        from io import BytesIO

        from openpyxl import Workbook

        from apps.finance.models import FinanceVendor

        FinanceVendor.objects.create(name='Rent Co', mobile='9000000000')
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Vendor Name', 'Mobile', 'PAN', 'MSME', 'TDS Applicable', 'TDS Section', 'TDS Rate (%)'])
        sheet.append(['rent co', 9111111111, 'aaafr1234k', 'Yes', 'Yes', '194I', 10])
        sheet.append(['New Vendor', '', '', 'No', 'No', '', ''])
        sheet.append(['Broken Vendor', '', 'BAD', 'No', 'No', '', ''])
        buffer = BytesIO()
        workbook.save(buffer)

        response = self._upload('/api/v1/finance/vendors/import/', buffer.getvalue())
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data['created'], 1)
        self.assertEqual(response.data['updated'], 1)
        self.assertEqual(len(response.data['errors']), 1)
        self.assertEqual(response.data['errors'][0]['row'], 4)
        vendor = FinanceVendor.objects.get(name='Rent Co')
        self.assertEqual(vendor.mobile, '+919111111111')
        self.assertEqual(vendor.pan, 'AAAFR1234K')
        self.assertTrue(vendor.msme_registered)
        self.assertTrue(vendor.tds_applicable)
        self.assertEqual(vendor.tds_rate, Decimal('10'))
        self.assertEqual(FinanceVendor.objects.count(), 2)

    def test_client_export_round_trip(self):
        from apps.finance.models import FinanceClient

        FinanceClient.objects.create(
            name='Round Trip Ltd', gstin='29ABCDE1234F1Z5', pan='ABCDE1234F', msme_registered=True,
            payment_terms_days=45,
        )
        response = self.client.get('/api/v1/finance/clients/export/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        content = b''.join(response.streaming_content) if response.streaming else response.content

        response = self.client.get('/api/v1/finance/clients/template/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        response = self._upload('/api/v1/finance/clients/import/', content)
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(
            {key: response.data[key] for key in ('created', 'updated', 'errors')},
            {'created': 0, 'updated': 1, 'errors': []},
        )
        self.assertIn('import_log_id', response.data)
        client = FinanceClient.objects.get(name='Round Trip Ltd')
        self.assertTrue(client.msme_registered)
        self.assertEqual(client.payment_terms_days, 45)


class WorkspaceFinanceTests(APITestCase):
    def setUp(self):
        self.finance_user = User.objects.create_user(
            email='workspace-finance@test.com',
            name='Workspace Finance',
            password='pass12345',
            role=UserRole.FINANCE,
        )
        self.client.force_authenticate(self.finance_user)

    def test_techpay_summary_ignores_techserv_rows(self):
        from apps.finance.models import PurchaseOrder, PurchaseOrderStatus

        PurchaseOrder.objects.create(
            po_number='PO-TECHSERV',
            customer_name='TechServ Customer',
            total_value=Decimal('100000'),
            status=PurchaseOrderStatus.ACTIVE,
        )
        techserv = self.client.get('/api/v1/finance/ops/summary/')
        self.assertEqual(techserv.status_code, status.HTTP_200_OK)
        self.assertEqual(techserv.data['po_metrics']['total_po_value'], 100000.0)

        techpay = self.client.get(
            '/api/v1/finance/ops/summary/',
            HTTP_X_WORKSPACE_COMPANY='techpay',
            HTTP_X_WORKSPACE_REGION='india',
        )
        self.assertEqual(techpay.status_code, status.HTTP_200_OK)
        self.assertEqual(techpay.data['po_metrics']['total_po_value'], 0.0)

    def test_created_invoice_stays_in_its_workspace(self):
        from apps.finance.models import Invoice

        created = self.client.post(
            '/api/v1/finance/invoices/',
            {
                'invoice_number': 'TP-1',
                'customer_name': 'TechPay Buyer',
                'amount': '1500.00',
                'invoice_date': '2026-04-01',
            },
            format='json',
            HTTP_X_WORKSPACE_COMPANY='techpay',
            HTTP_X_WORKSPACE_REGION='india',
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED, created.data)
        stored = Invoice.objects.get(invoice_number='TP-1')
        self.assertEqual(stored.workspace_company, 'techpay')
        self.assertEqual(stored.workspace_region, 'india')

        techserv_list = self.client.get('/api/v1/finance/invoices/')
        self.assertEqual(techserv_list.status_code, status.HTTP_200_OK)
        numbers = [row['invoice_number'] for row in techserv_list.data['results']]
        self.assertNotIn('TP-1', numbers)

        techpay_list = self.client.get(
            '/api/v1/finance/invoices/',
            HTTP_X_WORKSPACE_COMPANY='techpay',
            HTTP_X_WORKSPACE_REGION='india',
        )
        numbers = [row['invoice_number'] for row in techpay_list.data['results']]
        self.assertEqual(numbers, ['TP-1'])


class FinanceStatusTests(APITestCase):
    def setUp(self):
        from datetime import timedelta

        from apps.finance.models import Invoice, Payable

        self.users = {
            role: User.objects.create_user(
                email=f'status-{role.lower()}@test.com',
                name=f'{role.title()} User',
                password='pass12345',
                role=role,
            )
            for role in (UserRole.ADMIN, UserRole.FINANCE, UserRole.MANAGEMENT, UserRole.SALES)
        }
        self.today = date.today()
        self.invoice = Invoice.objects.create(
            invoice_number='ST-1',
            customer_name='Status Co',
            amount=Decimal('1000'),
            invoice_date=self.today - timedelta(days=40),
            due_date=self.today + timedelta(days=5),
        )
        self.payable = Payable.objects.create(
            vendor='Status Vendor',
            invoice_number='SP-1',
            total_amount=Decimal('500'),
            payable_amount=Decimal('500'),
            due_date=self.today - timedelta(days=3),
        )

    def _post(self, user, path, payload, **headers):
        self.client.force_authenticate(user)
        return self.client.post(path, payload, format='json', **headers)

    def test_new_rows_get_automatic_status(self):
        self.assertEqual(self.invoice.status, 'DUE')
        self.assertEqual(self.payable.status, 'OVERDUE')
        self.assertFalse(self.invoice.status_manual)

    def test_finance_and_admin_can_change_status_with_history(self):
        url = f'/api/v1/finance/invoices/{self.invoice.id}/status/'
        response = self._post(self.users[UserRole.FINANCE], url, {'status': 'received', 'note': 'NEFT 123'})
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data['status'], 'RECEIVED')
        self.assertTrue(response.data['status_manual'])
        self.assertEqual(Decimal(str(response.data['amount_received'])), Decimal('1000'))

        response = self._post(
            self.users[UserRole.ADMIN],
            f'/api/v1/finance/payables/{self.payable.id}/status/',
            {'status': 'PAID'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data['status'], 'PAID')

        self.client.force_authenticate(self.users[UserRole.SALES])
        history = self.client.get(f'/api/v1/finance/invoices/{self.invoice.id}/status-history/')
        self.assertEqual(history.status_code, status.HTTP_200_OK)
        self.assertEqual(len(history.data), 1)
        entry = history.data[0]
        self.assertEqual((entry['from_status'], entry['to_status']), ('DUE', 'RECEIVED'))
        self.assertEqual(entry['changed_by_name'], 'Finance User')
        self.assertEqual(entry['note'], 'NEFT 123')

    def test_other_roles_cannot_change_status(self):
        for role in (UserRole.MANAGEMENT, UserRole.SALES):
            response = self._post(
                self.users[role],
                f'/api/v1/finance/invoices/{self.invoice.id}/status/',
                {'status': 'RECEIVED'},
            )
            self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.invoice.refresh_from_db()
        self.assertEqual(self.invoice.status, 'DUE')
        self.assertFalse(self.invoice.status_changes.exists())

    def test_rejects_status_from_the_other_ledger(self):
        response = self._post(
            self.users[UserRole.FINANCE],
            f'/api/v1/finance/invoices/{self.invoice.id}/status/',
            {'status': 'PAID'},
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_crud_update_cannot_bypass_status_log(self):
        self.client.force_authenticate(self.users[UserRole.FINANCE])
        response = self.client.patch(
            f'/api/v1/finance/invoices/{self.invoice.id}/',
            {'status': 'RECEIVED', 'status_manual': True},
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.invoice.refresh_from_db()
        self.assertEqual(self.invoice.status, 'DUE')
        self.assertFalse(self.invoice.status_manual)

    def test_auto_refresh_moves_due_to_overdue_but_keeps_manual(self):
        from datetime import timedelta

        from apps.finance.models import Invoice
        from apps.finance.status import change_status, refresh_auto_statuses

        manual = Invoice.objects.create(
            invoice_number='ST-2',
            customer_name='Manual Co',
            amount=Decimal('200'),
            invoice_date=self.today,
            due_date=self.today + timedelta(days=5),
        )
        change_status(manual, 'DUE', self.users[UserRole.FINANCE], 'Promised next week')

        later = self.today + timedelta(days=10)
        refresh_auto_statuses(today=later)
        self.invoice.refresh_from_db()
        manual.refresh_from_db()
        self.assertEqual(self.invoice.status, 'OVERDUE')
        self.assertEqual(manual.status, 'DUE')

    def test_status_change_is_workspace_scoped(self):
        response = self._post(
            self.users[UserRole.FINANCE],
            f'/api/v1/finance/invoices/{self.invoice.id}/status/',
            {'status': 'RECEIVED'},
            HTTP_X_WORKSPACE_COMPANY='techpay',
            HTTP_X_WORKSPACE_REGION='india',
        )
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.client.force_authenticate(self.users[UserRole.FINANCE])
        history = self.client.get(
            f'/api/v1/finance/payables/{self.payable.id}/status-history/',
            HTTP_X_WORKSPACE_COMPANY='techpay',
            HTTP_X_WORKSPACE_REGION='india',
        )
        self.assertEqual(history.status_code, status.HTTP_404_NOT_FOUND)
